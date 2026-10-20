"""Construit Prototype_Maintenance_Preventive_V04.xlsx (V0.4 corrective et simplifiée).

Usage : python3 construire.py

1. génère les sources fictives DEMO_PDM.xlsx et DEMO_GMAO.xlsx (donnees_demo.py) ;
2. calcule l'état de démonstration avec la transposition Python des requêtes (regles.py) ;
3. écrit le classeur (openpyxl) : 5 onglets visibles, onglets techniques masqués ;
4. enregistre les résultats des formules (recalcul LibreOffice) ;
5. ajoute les requêtes Power Query intégrées, le tableau croisé et les segments (ooxml.py) ;
6. contrôle le fichier avec le validateur Open XML de Microsoft (outils/validation), s'il est installé.
"""

import datetime as dt
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ICI = Path(__file__).resolve().parent
sys.path.insert(0, str(ICI.parent))
import donnees_demo  # noqa: E402
import ooxml  # noqa: E402
import regles as R  # noqa: E402
from construire_prototype import (  # noqa: E402
    BLANC, FILET, GRIS_CLAIR, GRIS_MOYEN, JAUNE, POLICE, TEXTE, TEXTE_DISCRET, VERT, ajouter_logo, ecrire, fond,
    hauteur_texte, injecter_valeurs, police,
)

RACINE = ICI.parents[1]
V04 = RACINE / "V04"
SORTIE = V04 / "Prototype_Maintenance_Preventive_V04.xlsx"
LOGO = RACINE / "sources" / "logo_rvb.png"
VALIDATEUR = RACINE / "outils" / "validation"
ANNEE = 2027
ACTUALISATION = dt.datetime(2026, 10, 14, 15, 0)
JAUNE_PALE = "FFF4CC"

S_VISION, S_COUV, S_DETAIL, S_PARAM, S_NOTICE = (
    "01_VISION_GLOBALE", "02_COUVERTURE_OT", "03_DETAIL_OT", "04_PARAMETRES", "05_NOTICE")
S_CTL, S_ANO, S_PROG, S_RAPPRO, S_REF, S_CODE = (
    "90_CONTROLES", "91_ANOMALIES", "92_PROGRAMME", "93_RAPPROCHEMENT", "94_REFERENTIELS", "95_CODE_M")

REQUETES = [  # (nom de la requête, tableau chargé ou None, feuille)
    ("PDM_Normalise", "T_PDM", S_PROG),
    ("Vision_Globale", "T_Vision", S_VISION),
    ("GMAO_Lignes", None, None),
    ("OT_Detail", "T_OT_Detail", S_DETAIL),
    ("Rapprochement", "T_Rapprochement", S_RAPPRO),
    ("Controles", "T_Controles", S_CTL),
    ("Anomalies", "T_Anomalies", S_ANO),
]
ORDRE_ACTIVITES = ["PIS", "EC", "PILOTAGE", "FILTRE", "REV", "SURV DEC", "INSPECTION TSF", "INSPECTION TA", "ENDO"]

# Ligne d'en-tête des tableaux et ancres
L_VISION = 11
L_TCD = 12


def pied(ws):
    ws.oddFooter.left.text = "Prototype V0.4 – données 100 % fictives – la GMAO reste la source de référence"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "&A – page &P / &N"
    ws.oddFooter.right.size = 8


def impression(ws, zone, titres=None, paysage=True):
    ws.print_area = zone
    if titres:
        ws.print_title_rows = titres
    ws.page_setup.orientation = "landscape" if paysage else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    pied(ws)


def bandeau(ws, nb_col, titre, sous_titre, logo=True, etiquette="PROTOTYPE V0.4", col=None):
    for c in range(1, nb_col + 1):
        for r in (1, 2):
            ws.cell(r, c).fill = fond(VERT)
        ws.cell(3, c).fill = fond(JAUNE)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 3
    col = col or ("C" if logo else "A")
    ecrire(ws, f"{col}1", titre, font=police(18 if logo else 16, True, BLANC), fill=fond(VERT),
           alignment=Alignment(vertical="bottom", indent=0 if logo else 1))
    ecrire(ws, f"{col}2", sous_titre, font=police(9, couleur=BLANC), fill=fond(VERT),
           alignment=Alignment(vertical="center", indent=0 if logo else 1))
    der = get_column_letter(nb_col)
    if etiquette and der != col:
        ecrire(ws, f"{der}1", etiquette, font=police(9, True, JAUNE), fill=fond(VERT),
               alignment=Alignment(horizontal="right", vertical="bottom"))
        ecrire(ws, f"{der}2", "Données 100 % fictives", font=police(9, couleur=BLANC, italique=True), fill=fond(VERT),
               alignment=Alignment(horizontal="right", vertical="center"))
    if logo:
        ajouter_logo(ws, LOGO, hauteur_px=60)


def tableau(ws, nom, ligne, col, colonnes, lignes, formats=None, style="TableStyleMedium2", hauteur=None,
            police_corps=10, wrap=()):
    """Tableau Excel (chargé plus tard par une requête) : en-tête + données de démonstration."""
    formats = formats or {}
    for k, c in enumerate(colonnes):
        cell = ws.cell(ligne, col + k, c)
        cell.font = police(9.5, True, BLANC)
        cell.fill = fond(VERT)  # explicite : lisible même là où le style de tableau n'est pas appliqué
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, l in enumerate(lignes, start=ligne + 1):
        for k, c in enumerate(colonnes):
            v = l[c]
            cell = ws.cell(r, col + k, v)
            cell.font = police(police_corps, couleur=TEXTE)
            cell.alignment = Alignment(vertical="center", wrap_text=c in wrap)
            if c in formats:
                cell.number_format = formats[c]
            elif isinstance(v, dt.datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(v, dt.date):
                cell.number_format = "dd/mm/yyyy"
        if hauteur:
            ws.row_dimensions[r].height = hauteur
    fin = ligne + max(1, len(lignes))
    ref = f"{get_column_letter(col)}{ligne}:{get_column_letter(col + len(colonnes) - 1)}{fin}"
    t = Table(displayName=nom, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showFirstColumn=False, showLastColumn=False,
                                      showRowStripes=True, showColumnStripes=False)
    ws.add_table(t)
    ws.row_dimensions[ligne].height = 32
    return ref


def nommer(wb, nom, feuille, cellule):
    wb.defined_names[nom] = DefinedName(nom, attr_text=f"'{feuille}'!{cellule}")


def tuile(ws, debut, fin, ligne, libelle, formule, fmt="0"):
    for row, valeur, fnt, align in (
        (ligne, libelle, police(8, True, TEXTE_DISCRET), Alignment(horizontal="center", vertical="bottom", wrap_text=True)),
        (ligne + 1, formule, police(18, True, VERT), Alignment(horizontal="center", vertical="center")),
    ):
        if debut != fin:
            ws.merge_cells(f"{debut}{row}:{fin}{row}")
        ecrire(ws, f"{debut}{row}", valeur, font=fnt, fill=fond(GRIS_CLAIR), alignment=align, number_format=fmt)
        for c in range(openpyxl.utils.column_index_from_string(debut), openpyxl.utils.column_index_from_string(fin) + 1):
            ws.cell(row, c).fill = fond(GRIS_CLAIR)
            ws.cell(row, c).border = Border(left=Side(style="medium", color=BLANC), right=Side(style="medium", color=BLANC))


# ---------------------------------------------------------------- 01_VISION_GLOBALE

def feuille_vision(wb, ws, colonnes, vision):
    ws.sheet_properties.tabColor = JAUNE
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    larg = [9, 8, 12, 22, 15] + [18] * (len(colonnes) - 5)
    for i, w in enumerate(larg, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    nb = len(colonnes)
    bandeau(ws, nb, '="PROGRAMME DE MAINTENANCE PRÉVENTIVE "&MIN(T_PDM[Année])&"–"&MAX(T_PDM[Année])',
            "Une ligne = un poste : ce qui est prévu chaque année. Cliquez sur les segments pour filtrer.")
    for r in range(4, 10):
        ws.row_dimensions[r].height = 20
    # Trois informations seulement
    an = f'INDEX(T_Vision,0,MATCH(P_Annee&"",T_Vision[#Headers],0))'
    tuile(ws, "H", "I", 4, "Postes affichés", "=SUBTOTAL(103,T_Vision[Code poste])")
    tuile(ws, "J", "K", 4, '="Opérations préventives "&P_Annee',
          f'=IFERROR(SUMPRODUCT(SUBTOTAL(103,OFFSET(INDEX({an},1),ROW({an})-MIN(ROW({an})),0)),'
          f'--({an}<>""),LEN({an})-LEN(SUBSTITUTE({an},"+",""))+1),0)')
    tuile(ws, "L", "L", 4, "Année de référence", "=P_Annee")
    ws["J4"].comment = Comment("Somme des opérations prévues l'année de référence sur les postes affichés "
                               "(« PIS + EC SIMPLE » compte pour 2).", "Prototype", width=240, height=70)
    ws.merge_cells("H7:L9")
    ecrire(ws, "H7", "Cliquez sur un bouton de segment pour filtrer ; Ctrl + clic pour en choisir plusieurs ; "
           "l'entonnoir barré en haut du segment efface le filtre. Année de référence : 04_PARAMETRES.",
           font=police(8.5, couleur=TEXTE_DISCRET, italique=True),
           alignment=Alignment(wrap_text=True, vertical="top", horizontal="left", indent=1))
    ws.row_dimensions[10].height = 8
    tableau(ws, "T_Vision", L_VISION, 1, colonnes, vision, hauteur=26, wrap=colonnes[5:])
    for r in range(L_VISION + 1, L_VISION + 1 + len(vision)):
        ws.cell(r, 3).font = police(10, True, VERT)
    zone = f"A{L_VISION}:{get_column_letter(nb + 8)}{L_VISION + 3000}"
    ws.conditional_formatting.add(f"A{L_VISION}:{get_column_letter(nb + 8)}{L_VISION}", FormulaRule(
        formula=[f'AND(A${L_VISION}<>"",A${L_VISION}=P_Annee&"")'], fill=fond(JAUNE),
        font=Font(name=POLICE, bold=True, color=VERT)))
    ws.conditional_formatting.add(f"A{L_VISION + 1}:{get_column_letter(nb + 8)}{L_VISION + 3000}", FormulaRule(
        formula=[f'AND(A${L_VISION}<>"",A${L_VISION}=P_Annee&"")'], fill=fond(JAUNE_PALE)))
    ws.freeze_panes = f"A{L_VISION + 1}"
    impression(ws, f"A1:{get_column_letter(nb)}{L_VISION + len(vision)}", f"{L_VISION}:{L_VISION}")
    return larg


# ---------------------------------------------------------------- 02_COUVERTURE_OT

def ordre_items(valeurs, prefere=None):
    presents = list(dict.fromkeys(valeurs))
    debut = [v for v in (prefere or []) if v in presents]
    reste = sorted((v for v in presents if v not in debut and v is not None), key=str)
    return debut + reste + ([None] if None in presents else [])


def rendu_tcd(ot):
    """Valeurs affichées par les deux TCD (instantané, recalculé par Excel à l'ouverture)."""
    lignes_tcd = ["Année planifiée", "Site", "Équipe", "Code poste", "Désignation du poste"]
    partages = {c: ordre_items([o[c] for o in ot]) for c in lignes_tcd + ["Type de travail"]}
    partages["Activité atomique"] = ordre_items([o["Activité atomique"] for o in ot], ORDRE_ACTIVITES)
    activites = partages["Activité atomique"]
    cles = sorted({tuple(o[c] for c in lignes_tcd) for o in ot},
                  key=lambda k: tuple(partages[c].index(v) for c, v in zip(lignes_tcd, k)))
    comptes = {}
    for o in ot:
        k = (tuple(o[c] for c in lignes_tcd), o["Activité atomique"])
        comptes[k] = comptes.get(k, 0) + 1
    grille = [["Nb OT"] + [None] * 4 + ["Activité atomique"] + [None] * (len(activites) - 1),
              lignes_tcd + ["(vide)" if a is None else a for a in activites]]
    for k in cles:
        grille.append(list(k) + [comptes.get((k, a), 0) for a in activites])
    grille.append(["Total"] + [None] * 4 + [sum(comptes.get((k, a), 0) for k in cles) for a in activites])
    return partages, grille, lignes_tcd


def feuille_couverture(ws, ot):
    ws.sheet_properties.tabColor = JAUNE
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    partages, grille, lignes_tcd = rendu_tcd(ot)
    n_col = len(grille[1])
    larg = [10, 9, 9, 13, 24] + [11] * (n_col - 5)
    for i, w in enumerate(larg, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    bandeau(ws, max(n_col, 12), "COUVERTURE DES OT PRÉVENTIFS (GMAO)",
            "Nombre d'OT préventifs distincts (Type de travail PRE) par poste et par activité.")
    for r in range(4, 10):
        ws.row_dimensions[r].height = 20
    ws.merge_cells(f"A10:{get_column_letter(max(n_col, 12))}10")
    ecrire(ws, "A10", "Double-cliquez sur un nombre pour afficher le détail des OT concernés.",
           font=police(11, True, VERT), alignment=Alignment(vertical="center", indent=1), fill=fond(JAUNE_PALE))
    ws.row_dimensions[10].height = 24
    # Instantané des deux TCD (Excel les recalcule à l'ouverture depuis T_OT_Detail)
    kpi = [["Type de travail", "OT préventifs distincts"], ["PRE", sum(o["Compte OT"] for o in ot)]]
    for i, ligne in enumerate(kpi):
        for j, v in enumerate(ligne):
            c = ws.cell(4 + i, 7 + j, v)
            c.font = police(9 if i == 0 else 16, True, VERT)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, ligne in enumerate(grille):
        for j, v in enumerate(ligne):
            c = ws.cell(L_TCD + i, 1 + j, v)
            c.font = police(9.5, i < 2 or i == len(grille) - 1, VERT if i < 2 else TEXTE)
            if j >= 5 and i >= 2:
                c.alignment = Alignment(horizontal="center")
    fin = L_TCD + len(grille) - 1
    ws.conditional_formatting.add(f"F{L_TCD + 2}:AZ{L_TCD + 3000}", FormulaRule(
        formula=[f'F{L_TCD + 2}&""="0"'], font=Font(name=POLICE, color=GRIS_MOYEN)))
    ws.freeze_panes = f"A{L_TCD + 2}"
    impression(ws, f"A1:{get_column_letter(n_col)}{fin}", f"{L_TCD + 1}:{L_TCD + 1}")
    tcds = [
        {"nom": "TCD_Couverture", "ref": f"A{L_TCD}:{get_column_letter(n_col)}{fin}", "lignes": lignes_tcd,
         "colonne": "Activité atomique", "valeur": ("Ordre", "count", "Nb OT"), "premiere_ligne": 2, "total_bas": True},
        {"nom": "TCD_Total", "ref": "G4:H5", "lignes": ["Type de travail"], "colonne": None,
         "valeur": ("Compte OT", "sum", "OT préventifs distincts"), "premiere_ligne": 1, "total_bas": False,
         "filtres": ["Année planifiée", "Site", "Équipe"]},
    ]
    return partages, tcds, larg


# ---------------------------------------------------------------- 03, 90-95 : tables

def feuille_detail(ws, ot):
    ws.sheet_properties.tabColor = VERT
    larges = {"Ordre": 11, "Désignation": 40, "Désign.poste techn.": 26, "Désignation du poste": 20,
              "Contrôle qualité": 34, "Fichier source": 18, "Poste technique": 17, "Activité source": 20,
              "Importé le": 17, "Date début planifiée": 12, "Lignes export regroupées": 12}
    for i, c in enumerate(R.COLONNES_OT, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larges.get(c, 12)
    tableau(ws, "T_OT_Detail", 1, 1, R.COLONNES_OT, ot)
    ws.freeze_panes = "B2"
    impression(ws, f"A1:{get_column_letter(len(R.COLONNES_OT))}{len(ot) + 1}", "1:1")
    ws.page_setup.fitToWidth = 3


def feuille_table_simple(ws, nom, colonnes, lignes, larg=None):
    ws.sheet_properties.tabColor = GRIS_MOYEN
    for i, c in enumerate(colonnes, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (larg or {}).get(c, max(10, min(40, len(c) + 4)))
    tableau(ws, nom, 1, 1, colonnes, lignes, style="TableStyleLight9")
    ws.freeze_panes = "A2"
    ws.sheet_state = "hidden"
    pied(ws)


def feuille_referentiels(ws):
    ws.sheet_properties.tabColor = GRIS_MOYEN
    ws.sheet_state = "hidden"
    blocs = [
        ("A", "T_Ref_Sites", ["Site", "Statut"], [list(x) for x in R.SITES], [9, 12],
         "Sites (Groupe gestionnaires). « À valider » : conservé et signalé, pas mis en avant."),
        ("D", "T_Ref_Equipes", ["Équipe", "Statut"], [list(x) for x in R.EQUIPES], [9, 12],
         "Équipes (Poste travail princ.)."),
        ("G", "T_Ref_Activites", ["Activité", "Mot-clé dans Désignation (secours)", "Statut"],
         [[a, m, s] for a, m, s in R.ACTIVITES], [16, 22, 18],
         "Activités GMAO connues. Mot-clé : règle de secours, désactivée par défaut (04)."),
        ("K", "T_Ref_Operations", ["Opération PDM", "Libellé affiché", "Ordre"], [list(x) for x in R.OPERATIONS],
         [14, 14, 7], "Opérations du programme : libellé et ordre dans la vue 01."),
        ("O", "T_Colonnes_PDM", ["Champ", "Colonne dans le fichier PDM"], [list(x) for x in R.COLONNES_PDM], [16, 22],
         "Nom des colonnes du fichier PDM réel (à adapter, sans toucher aux requêtes)."),
        ("R", "T_Correspondance", ["Code poste PDM", "Code poste GMAO", "Commentaire"],
         [list(x) for x in R.CORRESPONDANCES], [14, 14, 30],
         "Exceptions seulement : par défaut Code Emprise = Code site."),
        ("V", "T_Regles_Comparaison", ["Opération programme", "Règle OT proposée", "Activités requises",
                                       "Activités exclues", "Statut"],
         [[a, b, c, d, "À VALIDER"] for a, b, c, d in R.REGLES], [14, 26, 14, 12, 11],
         "Future comparaison prévu / OT : NON appliquée tant que le métier ne l'a pas validée."),
    ]
    for col, nom, entetes, lignes, larg, aide in blocs:
        c0 = openpyxl.utils.column_index_from_string(col)
        ecrire(ws, f"{col}1", nom, font=police(11, True, VERT))
        ecrire(ws, f"{col}2", aide, font=police(8, couleur=TEXTE_DISCRET, italique=True))
        for k, w in enumerate(larg):
            ws.column_dimensions[get_column_letter(c0 + k)].width = w
        tableau(ws, nom, 4, c0, entetes, [dict(zip(entetes, l)) for l in lignes], style="TableStyleLight9")
    ws.freeze_panes = "A5"
    pied(ws)


def feuille_code(ws, requetes_code):
    ws.sheet_properties.tabColor = GRIS_MOYEN
    ws.sheet_state = "hidden"
    ws.column_dimensions["A"].width = 150
    ecrire(ws, "A1", "Code M des requêtes (déjà intégrées au classeur : Données › Requêtes et connexions). "
           "Copie de secours et de lecture.", font=police(11, True, VERT))
    r = 3
    for nom, code in requetes_code:
        ecrire(ws, f"A{r}", f"Requête « {nom} »", font=police(11, True, VERT))
        r += 1
        for ligne in code.splitlines():
            c = ws.cell(r, 1, ligne if ligne.strip() else None)
            c.font = Font(name="Consolas", size=9, color=TEXTE)
            r += 1
        r += 2


# ---------------------------------------------------------------- 04_PARAMETRES

def feuille_parametres(wb, ws):
    ws.sheet_properties.tabColor = VERT
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCD", [3, 40, 36, 62]):
        ws.column_dimensions[col].width = w
    bandeau(ws, 4, "PARAMÈTRES", "Quatre réglages, puis Données › Actualiser tout.", logo=False)
    lignes = [
        (5, "Fichier programme (PDM)", "DEMO_PDM.xlsx", "P_FichierPDM", True, "@",
         "Nom du fichier placé dans le même dossier que ce classeur, ou chemin complet (C:\\…\\PDM.xlsx)."),
        (6, "Fichier GMAO", "DEMO_GMAO.xlsx", "P_FichierGMAO", True, "@",
         "Export GMAO aux 24 colonnes habituelles (.xlsx ou .csv ;). Même règle que ci-dessus."),
        (7, "Année de référence", ANNEE, "P_Annee", True, "0",
         "Colonne mise en évidence dans 01 et compteur d'opérations."),
        (8, "Dernière actualisation", "=MAX(T_Controles[Actualisé le])", None, False, "dd/mm/yyyy hh:mm",
         "Écrite par les requêtes à chaque Données › Actualiser tout."),
    ]
    for ligne, lib, val, nom, saisie, fmt, aide in lignes:
        ecrire(ws, f"B{ligne}", lib, font=police(11, True, VERT), alignment=Alignment(vertical="center"))
        ecrire(ws, f"C{ligne}", val, font=police(11, saisie, TEXTE), number_format=fmt,
               fill=fond(BLANC) if saisie else fond(GRIS_CLAIR),
               alignment=Alignment(horizontal="left", vertical="center", indent=1),
               border=Border(left=FILET, right=FILET, top=FILET, bottom=FILET) if saisie else Border())
        ecrire(ws, f"D{ligne}", aide, font=police(8.5, couleur=TEXTE_DISCRET, italique=True),
               alignment=Alignment(vertical="center", wrap_text=True))
        ws.row_dimensions[ligne].height = 30
        if nom:
            nommer(wb, nom, S_PARAM, f"$C${ligne}")
    ws.row_dimensions[10].height = 34
    ecrire(ws, "B10", "Contrôles à vérifier", font=police(12, True, VERT), alignment=Alignment(vertical="center"))
    ecrire(ws, "C10", '=COUNTIF(T_Controles[Statut],"À VÉRIFIER")', font=police(20, True, VERT),
           alignment=Alignment(horizontal="left", vertical="center", indent=1), fill=fond(GRIS_CLAIR), number_format="0")
    ecrire(ws, "D10", "Détail dans l'onglet masqué 90_CONTROLES (clic droit sur un onglet › Afficher…).",
           font=police(8.5, couleur=TEXTE_DISCRET, italique=True), alignment=Alignment(vertical="center", wrap_text=True))
    ws.conditional_formatting.add("C10", FormulaRule(formula=["C10>0"], fill=fond(JAUNE)))
    ecrire(ws, "B13", "Réglages techniques (ne pas modifier sans raison)", font=police(9.5, True, TEXTE_DISCRET))
    ws["B13"].border = Border(bottom=Side(style="thin", color=GRIS_MOYEN))
    dossier = 'IFERROR(LEFT(CELL("filename",$A$1),FIND("[",CELL("filename",$A$1))-1),"")'
    chemin = lambda nom: (f'=IF(OR(ISNUMBER(SEARCH(":\\",{nom})),LEFT({nom},2)="\\\\",LEFT({nom},1)="/"),'
                          f'{nom},P_Dossier&{nom})')
    techniques = [
        (14, "Dossier de ce classeur (automatique)", f"={dossier}", "P_Dossier", False, "@"),
        (15, "Chemin lu pour le programme", chemin("P_FichierPDM"), "P_CheminPDM", False, "@"),
        (16, "Chemin lu pour la GMAO", chemin("P_FichierGMAO"), "P_CheminGMAO", False, "@"),
        (17, "Feuille du fichier PDM (vide = première)", None, "P_FeuillePDM", True, "@"),
        (18, "Feuille de l'export GMAO (vide = première)", None, "P_FeuilleGMAO", True, "@"),
        (19, "Longueur du code poste (repli sur Poste technique)", 9, "P_LongueurCodePoste", True, "0"),
        (20, "Reconstruire l'activité depuis la Désignation (secours)", "Non", "P_ActiviteSecours", True, "@"),
    ]
    for ligne, lib, val, nom, saisie, fmt in techniques:
        ecrire(ws, f"B{ligne}", lib, font=police(9, couleur=TEXTE_DISCRET), alignment=Alignment(vertical="center"))
        ecrire(ws, f"C{ligne}", val, font=police(9, couleur=TEXTE if saisie else TEXTE_DISCRET), number_format=fmt,
               fill=fond(BLANC) if saisie else fond(GRIS_CLAIR), alignment=Alignment(vertical="center", indent=1),
               border=Border(left=FILET, right=FILET, top=FILET, bottom=FILET) if saisie else Border())
        nommer(wb, nom, S_PARAM, f"$C${ligne}")
    dv = DataValidation(type="list", formula1='"Non,Oui"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C20")
    dv2 = DataValidation(type="whole", operator="between", formula1="2000", formula2="2100", showErrorMessage=True,
                         errorTitle="Année", error="Année sur 4 chiffres.")
    ws.add_data_validation(dv2)
    dv2.add("C7")
    ecrire(ws, "D20", "Non par défaut. « Oui » déduit l'activité de mots-clés de la Désignation quand la colonne "
           "Activité est vide : règle de secours à valider (94_REFERENTIELS).",
           font=police(8, couleur=TEXTE_DISCRET, italique=True), alignment=Alignment(wrap_text=True, vertical="center"))
    ws.row_dimensions[20].height = 30
    impression(ws, "A1:D20", paysage=False)


# ---------------------------------------------------------------- 05_NOTICE

NOTICE = [
    ("À quoi sert ce fichier", [
        "Voir, poste par poste, la maintenance préventive prévue cette année et les suivantes (01).",
        "En complément : vérifier que les OT préventifs correspondants existent dans la GMAO (02, 03).",
        "Uniquement du préventif : les OT dont le Type de travail n'est pas PRE sont exclus des vues.",
    ]),
    ("Utiliser", [
        "01_VISION_GLOBALE : cliquez sur SITE, ÉQUIPE, TYPOLOGIE (Ctrl + clic pour plusieurs). Une ligne = un poste.",
        "02_COUVERTURE_OT : choisissez SITE, ÉQUIPE, ANNÉE. Chaque nombre = OT distincts. Double-cliquez sur un "
        "nombre : Excel ouvre une nouvelle feuille avec les OT concernés (à supprimer après lecture).",
        "03_DETAIL_OT : tous les OT préventifs, filtres Excel standards (ex. filtrer un numéro d'Ordre).",
    ]),
    ("Mettre à jour (à chaque nouveau PDM ou export GMAO)", [
        "1. Remplacer DEMO_PDM.xlsx / DEMO_GMAO.xlsx par les nouveaux fichiers (ou saisir leur nom dans 04).",
        "2. Données › Actualiser tout. Les vues 01, 02 et 03, les contrôles et la date d'actualisation se mettent à jour.",
        "3. Lire « Contrôles à vérifier » dans 04.",
        "Première fois : si Excel demande les niveaux de confidentialité, choisir « Organisationnel » pour les deux "
        "fichiers et pour ce classeur.",
    ]),
    ("Règles appliquées", [
        "Site = Groupe gestionnaires ; Équipe = Poste travail princ. ; Année planifiée = année de Date début planifiée.",
        "Code poste GMAO = Code site ; s'il est vide, 9 premiers caractères du Poste technique (signalé).",
        "Code poste programme = Code Emprise ; s'il est vide, 9 premiers caractères du Code Installation (signalé).",
        "Activité découpée sur « + » (EC+PIS = 1 OT compté en EC et en PIS, une seule fois au total). Activité inconnue "
        "gardée et signalée « à valider ».",
        "Comptage = nombre d'Ordres distincts. Ordre vide : signalé, jamais compté. Lignes vides de fin d'export : ignorées.",
        "Rapprochement : Code Emprise = Code site, exceptions dans T_Correspondance (94). Aucun diagnostic prévu / OT.",
    ]),
    ("Onglets techniques (masqués)", [
        "90_CONTROLES, 91_ANOMALIES, 92_PROGRAMME, 93_RAPPROCHEMENT, 94_REFERENTIELS, 95_CODE_M. "
        "Clic droit sur un onglet › Afficher… pour les ouvrir (référentiels, correspondances, règles à valider).",
    ]),
]


def feuille_notice(ws):
    ws.sheet_properties.tabColor = VERT
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    ws.column_dimensions["C"].width = 22
    bandeau(ws, 3, "MODE D'EMPLOI", "L'essentiel en une page.", logo=False, col="B")
    r = 5
    for titre, lignes in NOTICE:
        ecrire(ws, f"B{r}", titre, font=police(12, True, VERT))
        ws[f"B{r}"].border = Border(bottom=Side(style="thin", color=JAUNE))
        ws.row_dimensions[r].height = 24
        r += 1
        for l in lignes:
            ecrire(ws, f"B{r}", l, font=police(10), alignment=Alignment(wrap_text=True, vertical="top"))
            ws.row_dimensions[r].height = hauteur_texte([(l, 110)], ligne=14, mini=18)
            r += 1
        r += 1
    impression(ws, f"A1:C{r}", paysage=False)


# ---------------------------------------------------------------- assemblage

def construire():
    donnees_demo.generer(V04)
    pdm = R.pdm_normalise(V04 / "DEMO_PDM.xlsx", fichier_source="DEMO_PDM.xlsx", importe_le=ACTUALISATION)
    colonnes_vision, vision = R.vision_globale(pdm)
    gmao = R.gmao_lignes(V04 / "DEMO_GMAO.xlsx", fichier_source="DEMO_GMAO.xlsx", importe_le=ACTUALISATION)
    ot = R.ot_detail(gmao, pdm)
    rappro = R.rapprochement(pdm, gmao)
    ctl = R.controles(gmao, pdm, rappro, ot, ACTUALISATION)
    ano = R.anomalies(gmao, pdm)

    wb = openpyxl.Workbook()
    wb._fonts = IndexedList([Font(name=POLICE, size=10, family=2)])
    wb.properties.title = "Prototype maintenance préventive V0.4"
    wb.properties.creator = "Prototype V0.4"
    ws = wb.active
    ws.title = S_VISION
    for nom in (S_COUV, S_DETAIL, S_PARAM, S_NOTICE, S_CTL, S_ANO, S_PROG, S_RAPPRO, S_REF, S_CODE):
        wb.create_sheet(nom)

    larg_vision = feuille_vision(wb, ws, colonnes_vision, vision)
    partages, tcds, larg_couv = feuille_couverture(wb[S_COUV], ot)
    feuille_detail(wb[S_DETAIL], ot)
    feuille_parametres(wb, wb[S_PARAM])
    feuille_notice(wb[S_NOTICE])
    feuille_table_simple(wb[S_CTL], "T_Controles", R.COLONNES_CTL, ctl,
                         {"Contrôle": 58, "Exemples": 44, "Actualisé le": 16, "N°": 6})
    feuille_table_simple(wb[S_ANO], "T_Anomalies", R.COLONNES_ANO, ano, {"Contrôle": 50, "Détail": 50})
    feuille_table_simple(wb[S_PROG], "T_PDM", R.COLONNES_PDM_NORM, pdm)
    feuille_table_simple(wb[S_RAPPRO], "T_Rapprochement", R.COLONNES_RAPPRO, rappro)
    feuille_referentiels(wb[S_REF])
    codes = [(nom, (V04 / "power-query" / f"{nom}.pq").read_text(encoding="utf-8")) for nom, _, _ in REQUETES]
    feuille_code(wb[S_CODE], codes)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    base = Path(tempfile.mkdtemp()) / SORTIE.name
    wb.save(base)

    # Résultats des formules en cache (aperçus sans calcul) ; chemins propres à la machine non figés
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(base)],
                       check=True, capture_output=True, timeout=300)
        figes = {(S_PARAM, "C14"): "", (S_PARAM, "C15"): "DEMO_PDM.xlsx", (S_PARAM, "C16"): "DEMO_GMAO.xlsx"}
        injecter_valeurs(base, Path(tmp) / base.name, figes)

    p = ooxml.Paquet(base)
    ooxml.ordonner_polices(p)
    ooxml.theme_charte(p, [VERT, JAUNE, "265B88", "5F8A8B", GRIS_MOYEN, GRIS_CLAIR], POLICE)
    colonnes_tables = {"T_PDM": R.COLONNES_PDM_NORM, "T_Vision": colonnes_vision, "T_OT_Detail": R.COLONNES_OT,
                       "T_Rapprochement": R.COLONNES_RAPPRO, "T_Controles": R.COLONNES_CTL, "T_Anomalies": R.COLONNES_ANO}
    effectifs = {"T_PDM": len(pdm), "T_Vision": len(vision), "T_OT_Detail": len(ot), "T_Rapprochement": len(rappro),
                 "T_Controles": len(ctl), "T_Anomalies": len(ano)}
    requetes = [{"nom": nom, "code": code, "tableau": t, "colonnes": colonnes_tables.get(t), "lignes": effectifs.get(t)}
                for (nom, t, _), (_, code) in zip(REQUETES, codes)]
    section = ooxml.ajouter_requetes(p, requetes, ooxml.LIAISONS)
    (V04 / "power-query" / "Section1.m").write_text(section, encoding="utf-8")
    ooxml.ajouter_tcd(p, S_COUV, "T_OT_Detail", R.COLONNES_OT, ot, partages, tcds)
    tables = p.tables()
    emu = lambda car: int(car * 7 * 9525 + 5 * 9525)  # largeur de colonne -> EMU (approximation Excel)
    marge = 4 * 9525
    ooxml.ajouter_segments(p, S_VISION, [
        {"nom": "SITE (vision)", "cache": "Segment_Site_Vision", "legende": "SITE", "source": "Site", "colonnes": 5,
         "table_id": tables["T_Vision"][1], "colonne_table": 1, "de": (0, marge, 3, marge), "a": (3, 0, 9, 0)},
        {"nom": "ÉQUIPE (vision)", "cache": "Segment_Equipe_Vision", "legende": "ÉQUIPE", "source": "Équipe",
         "colonnes": 1, "table_id": tables["T_Vision"][1], "colonne_table": 2, "de": (3, marge, 3, marge),
         "a": (4, 0, 9, 0)},
        {"nom": "TYPOLOGIE (vision)", "cache": "Segment_Typologie_Vision", "legende": "TYPOLOGIE",
         "source": "Typologie", "colonnes": 2, "table_id": tables["T_Vision"][1], "colonne_table": 5,
         "de": (4, marge, 3, marge), "a": (7, 0, 9, 0)},
    ], tableau=True, n_depart=1)
    ooxml.ajouter_segments(p, S_COUV, [
        {"nom": "SITE (OT)", "cache": "Segment_Site_OT", "legende": "SITE", "source": "Site", "colonnes": 5,
         "de": (0, marge, 3, marge), "a": (3, 0, 9, 0)},
        {"nom": "ÉQUIPE (OT)", "cache": "Segment_Equipe_OT", "legende": "ÉQUIPE", "source": "Équipe", "colonnes": 1,
         "de": (3, marge, 3, marge), "a": (4, 0, 9, 0)},
        {"nom": "ANNÉE (OT)", "cache": "Segment_Annee_OT", "legende": "ANNÉE", "source": "Année planifiée",
         "colonnes": 1, "de": (4, marge, 3, marge), "a": (5, 0, 9, 0)},
    ], tableau=False, n_depart=4, tcds=["TCD_Couverture", "TCD_Total"], champs_partages=partages)
    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    p.enregistrer(SORTIE)
    print(f"Écrit : {SORTIE} — {len(vision)} postes, {len(pdm)} lignes programme, {len(ot)} lignes OT x activité "
          f"({len({o['Ordre'] for o in ot})} OT), {sum(1 for c in ctl if c['Statut'] == 'À VÉRIFIER')} contrôles à vérifier")
    valider(SORTIE)
    return {"pdm": pdm, "vision": vision, "colonnes_vision": colonnes_vision, "gmao": gmao, "ot": ot,
            "rappro": rappro, "ctl": ctl, "ano": ano}


def valider(chemin):
    dll = VALIDATEUR / "out" / "Valid.dll"
    if not shutil.which("dotnet"):
        print("Validateur Open XML non disponible (dotnet absent).")
        return None
    if not dll.exists():
        subprocess.run(["dotnet", "build", "-c", "Release", "-o", "out"], cwd=VALIDATEUR, check=True,
                       capture_output=True, timeout=600)
    res = subprocess.run(["dotnet", str(dll), str(chemin)], capture_output=True, text=True, timeout=300)
    print(res.stdout.strip())
    return res.returncode == 0


if __name__ == "__main__":
    construire()
