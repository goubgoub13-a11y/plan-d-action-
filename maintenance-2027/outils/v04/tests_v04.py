"""Contrôles avant livraison de la V0.4 (liste du brief, point 28).

Lancement : python3 maintenance-2027/outils/v04/tests_v04.py
(après python3 maintenance-2027/outils/v04/construire.py)

Ce qui se vérifie hors Excel est vérifié ici : contenu du classeur, liens XML des segments et du
tableau croisé, règles (transposition Python des requêtes M), données 100 % fictives.
Le comportement réel dans Excel (clic sur un segment, double-clic, actualisation) ne l'est pas.
"""

import collections
import datetime as dt
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from lxml import etree

ICI = Path(__file__).resolve().parent
sys.path.insert(0, str(ICI))
import regles as R  # noqa: E402

V04 = ICI.parents[1] / "V04"
CLASSEUR = V04 / "Prototype_Maintenance_Preventive_V04.xlsx"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
      "x15": "http://schemas.microsoft.com/office/spreadsheetml/2010/11/main"}
RESULTATS = []


def verifier(libelle, condition, detail=""):
    RESULTATS.append((libelle, bool(condition), detail))
    print(("OK   " if condition else "ÉCHEC"), libelle, ("— " + detail) if detail else "")


def table(wb, nom):
    for ws in wb.worksheets:
        if nom in ws.tables:
            t = ws.tables[nom]
            lignes = list(ws[t.ref])
            entetes = [c.value for c in lignes[0]]
            return [dict(zip(entetes, [c.value for c in l])) for l in lignes[1:]], ws, t
    raise KeyError(nom)


def brut_gmao():
    ws = openpyxl.load_workbook(V04 / "DEMO_GMAO.xlsx", data_only=True).active
    lignes = [list(r) for r in ws.iter_rows(values_only=True)]
    entetes = [str(h).strip() for h in lignes[0]]
    return [dict(zip(entetes, l)) for l in lignes[1:]]


def brut_pdm():
    ws = openpyxl.load_workbook(V04 / "DEMO_PDM.xlsx", data_only=True).active
    lignes = [list(r) for r in ws.iter_rows(values_only=True)]
    return [dict(zip(lignes[0], l)) for l in lignes[1:] if any(v is not None for v in l)]


def main():
    wb = openpyxl.load_workbook(CLASSEUR, data_only=True)
    z = zipfile.ZipFile(CLASSEUR)
    vision, ws_vision, t_vision = table(wb, "T_Vision")
    ot, _, _ = table(wb, "T_OT_Detail")
    ctl = {c["Contrôle"]: c for c in table(wb, "T_Controles")[0]}
    gmao = brut_gmao()
    pdm = brut_pdm()

    # 1. Une ligne = un poste
    codes = [v["Code poste"] for v in vision]
    verifier("01 : une ligne = un poste", len(codes) == len(set(codes)) == len({(p["Code Emprise"] or p["Code Installation"][:9]) for p in pdm}),
             f"{len(codes)} lignes, {len(set(codes))} postes")
    verifier("01 : colonnes attendues",
             [c for c in vision[0]][:5] == ["Site", "Équipe", "Code poste", "Désignation du poste", "Typologie"]
             and [c for c in vision[0]][5:] == [str(a) for a in range(2026, 2033)])

    # 2. Plusieurs opérations la même année = une seule cellule « PIS + EC SIMPLE »
    ops = collections.defaultdict(set)
    for p in pdm:
        ops[(p["Code Emprise"] or p["Code Installation"][:9], p["Année"])].add(p["Opération"])
    par_code = {v["Code poste"]: v for v in vision}
    cas = [(c, a) for (c, a), o in ops.items() if o == {"PIS", "EC Simple"}]
    verifier("01 : PIS et EC Simple la même année -> « PIS + EC SIMPLE » dans une cellule",
             cas and all(par_code[c][str(a)] == "PIS + EC SIMPLE" for c, a in cas), f"{len(cas)} cas vérifiés")
    triple = [(c, a) for (c, a), o in ops.items() if o == {"PIS", "EC Filtre", "Révision"}]
    verifier("01 : trois opérations -> « PIS + EC FILTRE + REVISION »",
             triple and all(par_code[c][str(a)] == "PIS + EC FILTRE + REVISION" for c, a in triple), f"{len(triple)} cas")
    verifier("01 : toutes les cases correspondent au programme", all(
        (par_code[c][str(a)] or "") == " + ".join(sorted({R.OPERATIONS[[x[0] for x in R.OPERATIONS].index(op)][1] for op in o},
                                                         key=lambda l: ["PIS", "EC SIMPLE", "EC FILTRE", "REVISION"].index(l)))
        for (c, a), o in ops.items()))

    # 3. Segments reliés aux tableaux
    tables = {}
    for n in z.namelist():
        if n.startswith("xl/tables/table"):
            t = etree.fromstring(z.read(n))
            tables[t.get("name")] = t
    id_vision = tables["T_Vision"].get("id")
    cols_vision = [c.get("name") for c in tables["T_Vision"].find("m:tableColumns", NS)]
    segs = {}
    for n in z.namelist():
        if n.startswith("xl/slicerCaches/"):
            sc = etree.fromstring(z.read(n))
            segs[sc.get("name")] = sc
    tsc = {k: v.find(".//x15:tableSlicerCache", NS) for k, v in segs.items() if v.find(".//x15:tableSlicerCache", NS) is not None}
    verifier("Segments 01 reliés à T_Vision sur Site / Équipe / Typologie",
             sorted(cols_vision[int(e.get("column")) - 1] for e in tsc.values() if e.get("tableId") == id_vision)
             == ["Site", "Typologie", "Équipe"])
    psc = {k: v for k, v in segs.items() if k not in tsc}
    noms_tcd = {etree.fromstring(z.read(n)).get("name") for n in z.namelist() if re.fullmatch(r"xl/pivotTables/pivotTable\d+\.xml", n)}
    verifier("Segments 02 reliés aux deux TCD sur Site / Équipe / Année",
             sorted(v.get("sourceName") for v in psc.values()) == ["Année planifiée", "Site", "Équipe"]
             and all({p.get("name") for p in v.findall(".//x14:pivotTable", NS)} == noms_tcd for v in psc.values()))
    noms_slicers = {s.get("name") for n in z.namelist() if n.startswith("xl/slicers/")
                    for s in etree.fromstring(z.read(n))}
    noms_dessins = set(re.findall(rb'<sle:slicer [^>]*name="([^"]+)"', b"".join(
        z.read(n) for n in z.namelist() if n.startswith("xl/drawings/drawing"))))
    verifier("Chaque segment a sa forme dans la feuille", {n.encode() for n in noms_slicers} == noms_dessins,
             f"{len(noms_slicers)} segments")

    # 4. Seuls les OT PRE alimentent la couverture
    cache = etree.fromstring(z.read("xl/pivotCache/pivotCacheDefinition1.xml"))
    verifier("Couverture : uniquement Type de travail = PRE", {o["Type de travail"] for o in ot} == {"PRE"})
    cor = {str(g["Ordre"]) for g in gmao if g["Type de travail"] and g["Type de travail"] != "PRE"}
    verifier("Couverture : OT correctifs absents", cor and not cor & {o["Ordre"] for o in ot}, ", ".join(sorted(cor)))
    verifier("TCD alimenté par T_Optional_Detail".replace("Optional_", "OT_"),
             cache.find(".//m:worksheetSource", NS).get("name") == "T_OT_Detail" and int(cache.get("recordCount")) == len(ot))

    # 5. Ordres vides non comptés
    pre_vide = [g for g in gmao if g["Type de travail"] == "PRE" and g["Ordre"] is None and g["Désignation"]]
    attendu = len({str(g["Ordre"]) for g in gmao if g["Type de travail"] == "PRE" and g["Ordre"] is not None})
    kpi = wb["02_COUVERTURE_OT"]["H5"].value
    verifier("Ordre vide : signalé mais jamais compté", pre_vide and None not in {o["Ordre"] for o in ot}
             and ctl["Ordre vide"]["Nombre"] == len(pre_vide) and kpi == attendu,
             f"{len(pre_vide)} ligne(s) sans Ordre ; {kpi} OT distincts")

    # 6. Lignes techniques vides ignorées
    vides = [g for g in gmao if all(g[c] is None for c in R.STRUCTURANTS)]
    verifier("Lignes techniques vides ignorées (sans avertissement)",
             vides and ctl["Lignes techniques vides ignorées"]["Nombre"] == len(vides)
             and ctl["Lignes lues dans l'export (hors lignes vides)"]["Nombre"] == len(gmao) - len(vides),
             f"{len(vides)} lignes")

    # 7-8. Code site prioritaire et repli
    ecart = [g for g in gmao if g["Code site"] and g["Poste technique"] and g["Poste technique"][:9] != g["Code site"]]
    o = [x for x in ot if x["Ordre"] == str(ecart[0]["Ordre"])]
    verifier("Code site prioritaire sur GAUCHE(Poste technique;9)",
             o and o[0]["Code poste"] == ecart[0]["Code site"] and "Code site ≠" in (o[0]["Contrôle qualité"] or ""),
             f"{o[0]['Code poste']} (Poste technique {ecart[0]['Poste technique']})")
    sans = [g for g in gmao if g["Code site"] is None and g["Poste technique"]]
    o = [x for x in ot if x["Ordre"] == str(sans[0]["Ordre"])]
    verifier("Repli : Code site vide -> 9 premiers caractères du Poste technique",
             o and o[0]["Code poste"] == sans[0]["Poste technique"][:9] and o[0]["Source code poste"].startswith("Poste technique"))
    repli_pdm = {p["Code Installation"][:9] for p in pdm if p["Code Emprise"] is None}
    verifier("Repli programme : Code Emprise vide -> 9 premiers caractères du Code Installation",
             repli_pdm and repli_pdm <= set(codes), ", ".join(sorted(repli_pdm)))

    # 9. Comptage distinct (recalcul indépendant sur l'export brut)
    attendu = collections.defaultdict(set)
    for g in gmao:
        if g["Type de travail"] == "PRE" and g["Ordre"] is not None:
            code = g["Code site"] or g["Poste technique"][:9]
            # OT sans activité : compté une fois dans la colonne « (vide) » du TCD
            for a in (g["Activité"] or "(vide)").split("+"):
                a = a if a == "(vide)" else a.strip().upper()
                attendu[(g["Date début planifiée"].year, g["Groupe gestionnaires"], g["Poste travail princ."], code,
                         a)].add(g["Ordre"])
    ws = wb["02_COUVERTURE_OT"]
    entetes = [c.value for c in ws[13]]
    lu = {}
    for r in ws.iter_rows(min_row=14, values_only=True):
        if r[0] == "Total" or r[0] is None:
            break
        for k, a in enumerate(entetes[5:], start=5):
            if r[k]:
                lu[(r[0], r[1], r[2], r[3], a)] = r[k]
    attendu = {k: len(v) for k, v in attendu.items()}
    if lu != attendu:
        print("   écarts :", sorted(set(lu.items()) ^ set(attendu.items()), key=str)[:10])
    verifier("Comptage = nombre d'Ordres distincts (doublon compté une fois)",
             lu == attendu, f"{len(lu)} cases non nulles comparées")
    multi = collections.Counter(o["Ordre"] for o in ot)
    verifier("Un OT EC+PIS : 1 en EC, 1 en PIS, 1 seul OT au total",
             sum(o["Compte OT"] for o in ot) == len(multi) and any(n > 1 for n in multi.values()))

    # 10. Activité multiple éclatée
    g = next(x for x in gmao if x["Activité"] == "EC+PIS+PILOTAGE+FILTRE")
    verifier("« EC+PIS+PILOTAGE+FILTRE » -> 4 lignes atomiques",
             sorted(o["Activité atomique"] for o in ot if o["Ordre"] == str(g["Ordre"])) == ["EC", "FILTRE", "PILOTAGE", "PIS"])

    # 11. Nouveau site / activité / équipe / année : pas de rupture
    with tempfile.TemporaryDirectory() as tmp:
        src = openpyxl.load_workbook(V04 / "DEMO_GMAO.xlsx")
        ws = src.active
        ws.cell(2, 13, "NOUVEAU")           # Groupe gestionnaires
        ws.cell(3, 24, "EC+NOUVELLE ACT")   # Activité
        ws.cell(4, 12, "AA2")               # Poste travail princ. (à valider)
        ws.cell(5, 15, dt.datetime(2029, 5, 4))
        chemin = Path(tmp) / "GMAO_MODIFIE.xlsx"
        src.save(chemin)
        gm = R.gmao_lignes(chemin)
        o2 = R.ot_detail(gm)
        verifier("Nouveau site, activité, équipe, année : traités et signalés",
                 any(o["Site"] == "NOUVEAU" for o in o2) and any(o["Activité atomique"] == "NOUVELLE ACT" for o in o2)
                 and any(o["Année planifiée"] == 2029 for o in o2)
                 and any("Site inconnu" in r["Contrôles"] for r in gm)
                 and any("Activité inconnue (à valider)" in r["Contrôles"] for r in gm)
                 and any("Équipe à valider" in r["Contrôles"] for r in gm))
        p2 = openpyxl.load_workbook(V04 / "DEMO_PDM.xlsx")
        p2.active.append(["SET", "AA1", "G-I-AAA01", "G-I-AAA01-INS01", "Poste fictif Alpha", "DETENTE", 2033, "PIS"])
        p2.save(Path(tmp) / "PDM2.xlsx")
        cols, _ = R.vision_globale(R.pdm_normalise(Path(tmp) / "PDM2.xlsx"))
        verifier("Nouvelle année au programme : nouvelle colonne", cols[-1] == "2033")

    # 12-13. Données fictives, aucune donnée réelle ni chemin local
    textes = []
    for f in (V04 / "DEMO_GMAO.xlsx", V04 / "DEMO_PDM.xlsx", CLASSEUR):
        for wsx in openpyxl.load_workbook(f, data_only=True).worksheets:
            for row in wsx.iter_rows(values_only=True):
                textes += [str(v) for v in row if isinstance(v, str)]
    ordres = {str(g["Ordre"]) for g in gmao if g["Ordre"]}
    verifier("OT fictifs (910000001…)", all(re.fullmatch(r"9100000\d\d", o) for o in ordres))
    tous_codes = {g["Code site"] for g in gmao if g["Code site"]} | {p["Code Emprise"] for p in pdm if p["Code Emprise"]}
    verifier("Codes postes fictifs au format G-I-XXXnn", all(re.fullmatch(r"G-I-[A-Z]{3}\d{2}", c) for c in tous_codes))
    verifier("Désignations fictives", all(p["Désignation"].startswith("Poste fictif ") for p in pdm)
             and all(g["Désignation"] is None or " fictive " in g["Désignation"] for g in gmao))
    verifier("Sites et équipes : catégories réelles autorisées uniquement",
             {p["Site"] for p in pdm} | {g["Groupe gestionnaires"] for g in gmao if g["Groupe gestionnaires"]}
             <= {s for s, _ in R.SITES} and {p["Équipe"] for p in pdm} <= {"AA1", "AA3", "AA4"})
    interdits = ["SITE0", "TECH-0", "G-I-A1100", "INST-0", "/home/", "/tmp/", "file:///"]
    blob = b"".join(z.read(n) for n in z.namelist())
    verifier("Aucune donnée des versions précédentes ni chemin local dans le livrable",
             not any(i in t for t in textes for i in interdits) and not any(i.encode() in blob for i in interdits[4:]))

    print(f"\n{sum(ok for _, ok, _ in RESULTATS)} / {len(RESULTATS)} contrôles réussis")
    return all(ok for _, ok, _ in RESULTATS)


if __name__ == "__main__":
    sys.exit(0 if main() else 1)
