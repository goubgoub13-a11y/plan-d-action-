"""Génère les deux fichiers sources de démonstration de la V0.4 (100 % fictifs).

- DEMO_PDM.xlsx  : programme préventif pluriannuel, une ligne par poste x année x opération ;
- DEMO_GMAO.xlsx : export GMAO aux 24 en-têtes réels (dont « Typologie » avec espace final).

Seules les catégories fonctionnelles sont réelles (sites, équipes, activités, types du
programme, typologies). Codes postes, désignations, numéros d'OT, dates et volumes sont
inventés. L'export contient volontairement quelques cas à contrôler (OT correctifs, Ordre
vide, doublon, Code site vide ou incohérent, activité issue d'une formule, lignes
techniques vides en fin de tableau).
"""

import datetime as dt
import random
from pathlib import Path

import openpyxl

SITES = ["SET", "SPS", "RNA", "VRS", "LRO", "ALZ", "VDY", "DEI", "SHL", "SVI"]
EQUIPES = ["AA1", "AA3", "AA4"]
TYPOLOGIES = ["DETENTE", "LIVRAISON", "SECTIONNEMENT", "COUPURE", "COMPTAGE", "PIQUAGE"]
GRECS = ["Alpha", "Beta", "Gamma", "Delta", "Epsilon", "Zeta", "Eta", "Theta", "Iota", "Kappa", "Lambda", "Mu",
         "Nu", "Xi", "Omicron", "Pi", "Rho", "Sigma", "Tau", "Upsilon", "Phi", "Chi", "Psi", "Omega"]
ANNEES = list(range(2026, 2033))
ANNEES_GMAO = [2026, 2027]

ENTETES_GMAO = ["Créé le", "Ordre", "Ordre supérieur", "Type d'ordre", "Type de travail", "Avis", "Désignation",
                "Statut utilisateur", "Statut système", "Saisi par", "Total coûts réels", "Poste travail princ.",
                "Groupe gestionnaires", "Secteur exploitation", "Date début planifiée", "Code site",
                "Poste technique", "Désign.poste techn.", "Équipement", "Groupe de gamme principale",
                "Compteur de groupe de gamme principale", "Date modif. fiche ordre", "Typologie ", "Activité"]
ENTETES_PDM = ["Site", "Équipe", "Code Emprise", "Code Installation", "Désignation", "Typologie", "Année",
               "Opération"]

# Cas particuliers de la démonstration
POSTE_SANS_CODE_EMPRISE = (4, 17)        # code poste repris des 9 premiers caractères de Code Installation
POSTE_CORRESPONDANCE_MANUELLE = 12       # code GMAO différent du code PDM -> table de correspondance
CODE_GMAO_MANUEL = "G-I-MMN13"
POSTE_SANS_OT = 25                       # aucun OT dans l'export -> « Non rapproché »


def code_poste(i):
    lettres = chr(65 + i % 26) * 3 if i < 26 else chr(65 + i - 26) * 2 + chr(66 + i - 26)
    return f"G-I-{lettres}{i + 1:02d}"


def postes():
    extra = [("DEI", "AA1"), ("VRS", "AA3"), ("SET", "AA3"), ("LRO", "AA4"), ("RNA", "AA1"), ("SVI", "AA4")]
    paires = [(s, e) for s in SITES for e in EQUIPES] + extra
    rnd = random.Random(2027)
    liste = []
    for i, (site, equipe) in enumerate(paires):
        nom = GRECS[i % 24] + ("" if i < 24 else " II")
        liste.append({
            "i": i, "site": site, "equipe": equipe, "code": code_poste(i), "designation": f"Poste fictif {nom}",
            "typologie": TYPOLOGIES[rnd.randrange(len(TYPOLOGIES))],
            "pis_annuel": rnd.random() < 0.7, "pis_phase": rnd.randrange(2),
            "ec_periode": rnd.choice([2, 3]), "ec_phase": rnd.randrange(3), "ec_filtre": rnd.random() < 0.5,
            "rev_phase": rnd.randrange(6), "rev_seule": rnd.random() < 0.4,
        })
    return liste


def programme(p):
    """Opérations prévues par année pour un poste : {année: [opérations]}."""
    plan = {}
    for k, annee in enumerate(ANNEES):
        ops = []
        revision = (k - p["rev_phase"]) % 6 == 0
        if p["pis_annuel"] or k % 2 == p["pis_phase"]:
            if not (revision and p["rev_seule"]):
                ops.append("PIS")
        if (k - p["ec_phase"]) % p["ec_periode"] == 0 and not (revision and p["rev_seule"]):
            occurrence = (k - p["ec_phase"]) // p["ec_periode"]
            ops.append("EC Filtre" if p["ec_filtre"] and occurrence % 2 == 0 else "EC Simple")
        if revision:
            ops.append("Révision")
        if ops:
            plan[annee] = ops
    return plan


def ecrire_pdm(chemin, liste):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Programme"
    ws.append(ENTETES_PDM)
    for p in liste:
        for annee, ops in programme(p).items():
            for op in ops:
                emprise = None if p["i"] in POSTE_SANS_CODE_EMPRISE else p["code"]
                ws.append([p["site"], p["equipe"], emprise, f"{p['code']}-INS{p['i'] % 7 + 1:02d}",
                           p["designation"], p["typologie"], annee, op])
    for _ in range(2):  # lignes techniques vides en fin de tableau
        ws.append([None] * len(ENTETES_PDM))
    for col, w in zip("ABCDEFGH", [8, 8, 13, 20, 24, 15, 8, 12]):
        ws.column_dimensions[col].width = w
    wb.save(chemin)


ACTIVITES_DE = {"PIS": ["PIS"], "EC Simple": ["EC"], "EC Filtre": ["EC", "FILTRE"], "Révision": ["REV"]}
ORDRE_ACT = ["EC", "PIS", "PILOTAGE", "REV", "FILTRE"]
SUFFIXES = ["REG01", "VAN01", "FILT01", "SEC01"]


def ecrire_gmao(chemin, liste):
    rnd = random.Random(910)
    ordre = iter(range(910000001, 910009999))
    lignes = []

    def ot(p, annee, activites, mois=None, type_travail="PRE", code_site=None, poste_tech=None, ordre_ot=None,
           activite=None):
        acts = "+".join(sorted(activites, key=lambda a: ORDRE_ACT.index(a) if a in ORDRE_ACT else 9))
        mois = mois or rnd.randint(1, 11)
        debut = dt.datetime(annee, mois, rnd.randint(1, 28))
        code = code_site if code_site is not None else (CODE_GMAO_MANUEL if p["i"] == POSTE_CORRESPONDANCE_MANUELLE
                                                         else p["code"])
        suffixe = SUFFIXES[p["i"] % 4]
        pt = poste_tech if poste_tech is not None else f"{code or p['code']}-{suffixe}"
        lignes.append([
            debut - dt.timedelta(days=rnd.randint(40, 90)), ordre_ot if ordre_ot is not None else next(ordre), None,
            "ZPR2" if type_travail == "PRE" else "ZCO1", type_travail, None,
            f"Maintenance {'préventive' if type_travail == 'PRE' else 'corrective'} fictive {acts} - poste {pt[:9]}",
            rnd.choice(["ACTI", "ACTI", "TERM"]), "STAT_SYS_A", rnd.choice(["USER01", "USER02", "USER03"]),
            round(rnd.uniform(150, 1500), 2), p["equipe"], p["site"], p["equipe"], debut, code, pt,
            f"Équipement fictif {pt}", f"EQ-{pt[4:9]}-{rnd.randint(1, 9):02d}", "GG-PREV", rnd.randint(1, 3),
            debut - dt.timedelta(days=rnd.randint(5, 30)), p["typologie"], activite if activite is not None else acts])
        return lignes[-1]

    for p in liste:
        if p["i"] == POSTE_SANS_OT:
            continue
        plan = programme(p)
        for annee in ANNEES_GMAO:
            ops = plan.get(annee, [])
            if not ops or rnd.random() < 0.08:  # quelques opérations prévues sans OT
                continue
            activites = [a for op in ops for a in ACTIVITES_DE[op]]
            if "FILTRE" in activites and rnd.random() < 0.5:
                activites.insert(1, "PILOTAGE")
            if "PIS" in activites and len(activites) > 1 and rnd.random() < 0.25:
                ot(p, annee, ["PIS"])                    # PIS dans un OT séparé
                activites = [a for a in activites if a != "PIS"]
            if activites == ["PIS"] and rnd.random() < 0.2:
                ot(p, annee, ["PIS"], mois=3)            # deux visites PIS la même année
                ot(p, annee, ["PIS"], mois=9)
                continue
            ot(p, annee, activites)

    par_code = {p["code"]: p for p in liste}
    for code, act in [("G-I-AAA01", ["SURV DEC"]), ("G-I-HHH08", ["SURV DEC"]), ("G-I-DDD04", ["INSPECTION TSF"]),
                      ("G-I-KKK11", ["INSPECTION TA"]), ("G-I-VVV22", ["ENDO"])]:
        ot(par_code[code], 2027, act)
    # Cas à contrôler
    ot(par_code["G-I-BBB02"], 2026, ["EC"], type_travail="COR")
    ot(par_code["G-I-QQQ17"], 2027, ["PIS"], type_travail="COR")
    vide = ot(par_code["G-I-CCC03"], 2027, ["PIS"])
    vide[1] = None                                             # Ordre vide
    lignes.append(list(ot(par_code["G-I-EEE05"], 2026, ["EC", "PIS"])))   # doublon exact d'un OT
    ot(par_code["G-I-FFF06"], 2027, ["PIS"], code_site="")                 # Code site vide -> repli
    ot(par_code["G-I-GGG07"], 2027, ["EC"], poste_tech="G-I-GGX07-VAN01")  # Code site ≠ Poste technique
    ot(par_code["G-I-JJJ10"], 2027, ["PIS", "CONTROLE VISUEL"])            # activité inconnue
    formule = ot(par_code["G-I-LLL12"], 2027, ["PIS"])
    n_formule = len(lignes) + 1                                            # Activité issue d'une formule
    formule[23] = f'=IF(ISNUMBER(SEARCH("PIS",G{n_formule})),"PIS","")'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(ENTETES_GMAO)
    for l in lignes:
        ws.append([None if v == "" else v for v in l])
    for k in range(3):  # lignes techniques vides : seule la formule d'Activité subsiste
        r = len(lignes) + 2 + k
        ws.cell(r, 24, f'=IF(ISNUMBER(SEARCH("PIS",G{r})),"PIS","")')
    for r in range(2, ws.max_row + 1):
        for c in (1, 15, 22):
            ws.cell(r, c).number_format = "dd/mm/yyyy"
    wb.save(chemin)
    return len(lignes)


def generer(dossier):
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    liste = postes()
    ecrire_pdm(dossier / "DEMO_PDM.xlsx", liste)
    n = ecrire_gmao(dossier / "DEMO_GMAO.xlsx", liste)
    return liste, n


if __name__ == "__main__":
    liste, n = generer(Path(__file__).resolve().parents[2] / "V04")
    print(f"{len(liste)} postes fictifs, {n} lignes GMAO (+3 lignes techniques vides)")
