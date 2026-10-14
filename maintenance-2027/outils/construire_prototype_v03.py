"""Construit Prototype_Maintenance_Preventive_V03.xlsx (V0.3 : programme préventif + couche OT GMAO).

Usage :
    python3 construire_prototype_v03.py [--sortie FICHIER]

Le classeur sépare deux couches :
- programme préventif (cœur du produit) : 01_PLAN_PREVENTIF, 03_DETAIL_PROGRAMME ;
- OT préventifs GMAO (complément) : 02_COUVERTURE_OT, 04_DETAIL_OT et les contrôles de 08.

Les zones OT contiennent un instantané de démonstration calculé par gmao_ot.py, avec
les mêmes règles et le même gabarit que les requêtes Power Query de V03/power-query.
Une fois ces requêtes installées (09_NOTICE), « Actualiser tout » les remplace.
"""

import argparse
import datetime as dt
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import gmao_ot
from construire_prototype import (
    BLANC, FILET, GRIS_CLAIR, GRIS_MOYEN, JAUNE, POLICE, SEPARATEUR_BLANC, TEXTE, TEXTE_DISCRET, VERT,
    ajouter_logo, ajouter_tableau, bandeau, ecrire, en_tete_tableau, fond, hauteur_texte, injecter_valeurs,
    largeurs, lire_table, mise_en_page, police,
)

ICI = Path(__file__).resolve().parent
RACINE = ICI.parent
SOURCES = RACINE / "sources"
SORTIE = RACINE / "V03" / "Prototype_Maintenance_Preventive_V03.xlsx"
CODE_M = RACINE / "V03" / "power-query"

VERSION = "V0.3"
ANNEE = 2027
FICHIER_GMAO = "MODELE_GMAO_PREV_ANONYMISE.csv"
INSTANTANE_LE = dt.datetime(2026, 10, 14, 12, 11)  # date du fichier modèle reçu
IMPORT_PROGRAMME_LE = dt.datetime(2026, 10, 13)
ACTIVITES = ["PIS", "EC", "PILOTAGE", "FILTRE", "REV"]  # ordre du brief, toutes observées dans le modèle
HYPOTHESES_ACTIVITES = {"FILTRE": "TYPE_C « Contrôle filtre » ?", "REV": "TYPE_D « Révision » ?"}

S_PLAN = "01_PLAN_PREVENTIF"
S_COUV = "02_COUVERTURE_OT"
S_PROG = "03_DETAIL_PROGRAMME"
S_OT = "04_DETAIL_OT"
S_REF = "05_REFERENTIEL"
S_AVIS = "06_AVIS_UTILISATEURS"
S_POINTS = "07_POINTS_A_VALIDER"
S_PARAM = "08_PARAMETRES"
S_NOTICE = "09_NOTICE"
S_CODE = "10_CODE_M"

# Zones d'atterrissage des requêtes Power Query (cellule en haut à gauche du tableau chargé)
ANCRE_COUV = (8, 1)      # 02_COUVERTURE_OT!A8
ANCRE_OT = (1, 1)        # 04_DETAIL_OT!A1
ANCRE_CTL = (20, 1)      # 08_PARAMETRES!A20
MAX_LIGNES = 200000
MAX_COL = "AZ"

PIED = f"Prototype {VERSION} – données fictives – la GMAO reste la source de référence"


def pied(ws):
    ws.oddFooter.left.text = PIED


def plage_couv(col_debut="A", col_fin=MAX_COL):
    l = ANCRE_COUV[0] + 1
    return f"'{S_COUV}'!${col_debut}${l}:${col_fin}${MAX_LIGNES}"


ENTETE_COUV = f"'{S_COUV}'!$A${ANCRE_COUV[0]}:${MAX_COL}${ANCRE_COUV[0]}"
ANNEES_COUV = f"'{S_COUV}'!$A${ANCRE_COUV[0] + 1}:$A${MAX_LIGNES}"
CTL_STATUTS = f"'{S_PARAM}'!$G${ANCRE_CTL[0] + 1}:$G${ANCRE_CTL[0] + 60}"


def colonne_couv(nom):
    """Colonne de la couverture repérée par son en-tête (les activités varient d'un export à l'autre)."""
    return f"INDEX({plage_couv()},0,MATCH(\"{nom}\",{ENTETE_COUV},0))"


def valeur_detail(nom):
    """Première valeur d'une colonne de 04_DETAIL_OT, repérée par son en-tête."""
    return (f"IFERROR(INDEX('{S_OT}'!$A$2:${MAX_COL}$2,MATCH(\"{nom}\",'{S_OT}'!$A$1:${MAX_COL}$1,0)),\"\")")


def tuiles(ws, ligne, definitions, groupes=()):
    """definitions : (col_debut, col_fin, libellé, formule). Libellé ligne n, valeur ligne n+1."""
    from openpyxl.utils import column_index_from_string as idx
    for debut, fin, libelle, formule in definitions:
        for row, valeur, fnt, align in (
            (ligne, libelle, police(8, True, TEXTE_DISCRET),
             Alignment(horizontal="center", vertical="bottom", wrap_text=True)),
            (ligne + 1, formule, police(16, True, VERT), Alignment(horizontal="center", vertical="center")),
        ):
            if debut != fin:
                ws.merge_cells(f"{debut}{row}:{fin}{row}")
            ecrire(ws, f"{debut}{row}", valeur, font=fnt, fill=fond(GRIS_CLAIR), alignment=align,
                   number_format="0")
            for c in range(idx(debut), idx(fin) + 1):
                cell = ws.cell(row, c)
                cell.fill = fond(GRIS_CLAIR)
                droite = Side(style="thick", color=BLANC) if get_column_letter(c) in groupes else SEPARATEUR_BLANC
                cell.border = Border(left=SEPARATEUR_BLANC, right=droite)
    ws.row_dimensions[ligne].height = 24
    ws.row_dimensions[ligne + 1].height = 26


def entete_vue(ws, nb_col, titre, sous_titre, logo=None, etiquette=f"PROTOTYPE {VERSION}",
               mention="Données 100 % fictives"):
    der = get_column_letter(nb_col)
    bandeau(ws, nb_col, None, None)
    col_titre = "C" if logo else "A"
    ecrire(ws, f"{col_titre}1", titre, font=police(18 if logo else 16, True, BLANC), fill=fond(VERT),
           alignment=Alignment(vertical="bottom", indent=0 if logo else 1))
    ecrire(ws, f"{col_titre}2", sous_titre, font=police(9, couleur=BLANC), fill=fond(VERT),
           alignment=Alignment(vertical="center", indent=0 if logo else 1))
    if der != col_titre:
        if etiquette:
            ecrire(ws, f"{der}1", etiquette, font=police(9, True, JAUNE), fill=fond(VERT),
                   alignment=Alignment(horizontal="right", vertical="bottom"))
        if mention:
            ecrire(ws, f"{der}2", mention, font=police(9, couleur=BLANC, italique=True), fill=fond(VERT),
                   alignment=Alignment(horizontal="right", vertical="center"))
    if logo:
        ajouter_logo(ws, logo, hauteur_px=60)


def ligne_info(ws, ligne, cellules):
    """cellules : (ref, valeur, style) ; style = 'etiquette' | 'valeur' | 'aide'."""
    styles = {
        "etiquette": (police(8, couleur=TEXTE_DISCRET), Alignment(horizontal="right", vertical="center")),
        "valeur": (police(10, True, VERT), Alignment(horizontal="center", vertical="center")),
        "texte": (police(8, True, VERT), Alignment(horizontal="left", vertical="center", indent=1)),
        "aide": (police(8, couleur=TEXTE_DISCRET, italique=True), Alignment(horizontal="right", vertical="center")),
    }
    for ref, valeur, style in cellules:
        f, a = styles[style]
        ecrire(ws, ref, valeur, font=f, alignment=a)
    ws.row_dimensions[ligne].height = 20


def regle(formule, **style):
    return FormulaRule(formula=[formule], **style)


JAUNE_TEXTE = dict(fill=fond(JAUNE), font=Font(name=POLICE, bold=True, color=VERT))


# ---------------------------------------------------------------- 01_PLAN_PREVENTIF

def feuille_plan(ws, lignes_prog, cols_prog, ref_lignes, logo):
    i = {n: k for k, n in enumerate(cols_prog)}
    ordre_type = {l[0]: k for k, l in enumerate(ref_lignes)}
    annees = list(range(min(l[i["Année"]] for l in lignes_prog), max(l[i["Année"]] for l in lignes_prog) + 1))
    cles = sorted({(l[i["Périmètre"]], l[i["Groupe"]], l[i["Code installation"]], l[i["Installation"]],
                    l[i["Type opération source"]], l[i["Périodicité"]]) for l in lignes_prog},
                  key=lambda k: (k[0], k[1], k[2], ordre_type.get(k[4], 99), k[5]))

    colonnes = (["Périmètre", "Groupe", "Code installation", "Installation", "Opération", "Périodicité"]
                + [str(a) for a in annees] + ["Code type source", "Statut référentiel", "Point de vigilance"])
    larg = [10, 9, 12, 15, 19, 10] + [7] * len(annees) + [11, 15, 24]
    nb = len(colonnes)
    c_an0, c_an1 = 7, 6 + len(annees)
    L_AN0, L_AN1 = get_column_letter(c_an0), get_column_letter(c_an1)
    C_TYPE, C_STATUT, C_VIG = (get_column_letter(c_an1 + k) for k in (1, 2, 3))
    der = get_column_letter(nb)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = JAUNE
    largeurs(ws, larg)
    entete_vue(ws, nb, "Plan préventif pluriannuel",
               "Site → Équipe → Poste → Année. Programme préventif par poste ; la GMAO reste la source de référence.",
               logo=logo)

    t = "T_Plan"
    tuiles(ws, 4, [
        ("A", "B", '="Opérations prévues "&P_Annee', "=COUNTIF(T_Programme[Année],P_Annee)"),
        ("C", "D", '="Postes concernés "&P_Annee',
         "=SUMPRODUCT((T_Programme[Année]=P_Annee)/COUNTIFS(T_Programme[Code installation],"
         "T_Programme[Code installation],T_Programme[Année],T_Programme[Année]))"),
        ("E", "E", "Opérations au programme (toutes années)", "=COUNTA(T_Programme[Code installation])"),
        ("F", "G", "Postes au programme",
         '=SUMPRODUCT((T_Programme[Code installation]<>"")/COUNTIF(T_Programme[Code installation],'
         'T_Programme[Code installation]&""))'),
        ("H", "J", "Années couvertes", '=MIN(T_Programme[Année])&" – "&MAX(T_Programme[Année])'),
        ("K", "M", '="OT GMAO "&P_Annee&" (complément)"',
         f"=IFERROR(SUMIFS({colonne_couv('OT distincts (toutes activités)')},{ANNEES_COUV},P_Annee),0)"),
        (C_TYPE, C_STATUT, "Statut à confirmer",
         f'=COUNTIF({t}[Statut référentiel],"À confirmer")+COUNTIF({t}[Statut référentiel],"Absent du référentiel")'),
        (C_VIG, C_VIG, "Points de vigilance", f'=COUNTIF({t}[Point de vigilance],"?*")'),
    ], groupes={"D", "J", "M"})
    ws["K4"].comment = Comment("Nombre d'OT préventifs distincts trouvés dans l'extraction GMAO pour l'année "
                               "d'analyse (onglet 02). Simple repère : aucun rapprochement avec le programme "
                               "n'est fait tant que la clé de correspondance n'est pas validée.", "Prototype",
                               width=280, height=100)

    premiere, derniere = 8, 7 + len(cles)
    nb_lignes_prog = len(lignes_prog)
    ligne_info(ws, 6, [
        ("A6", "Année d'analyse", "etiquette"),
        ("B6", "=P_Annee", "valeur"),
        ("C6", f'=IF(SUM({t}[[{annees[0]}]:[{annees[-1]}]])=COUNTA(T_Programme[Année]),'
               f'"Contrôle programme : OK ("&COUNTA(T_Programme[Année])&" lignes)",'
               f'"Contrôle programme : VUE À ACTUALISER (source "&COUNTA(T_Programme[Année])&'
               f'" / vue "&SUM({t}[[{annees[0]}]:[{annees[-1]}]])&")")', "texte"),
        (f"{der}6", "● opération prévue  ·  ●×2 lignes identiques  ·  colonne jaune = année d'analyse (08)", "aide"),
    ])
    ws.merge_cells("C6:H6")
    ws["B6"].number_format = "0"

    en_tete_tableau(ws, 7, colonnes, hauteur=32)
    match = f"MATCH(${C_TYPE}{{r}},T_Referentiel[Valeur source],0)"
    for k, (perim, groupe, code, inst, typ, perio) in enumerate(cles):
        r = premiere + k
        valeurs = [perim, groupe, code, inst,
                   f'=IFERROR(INDEX(T_Referentiel[Libellé affiché],{match}),"À confirmer")', perio]
        for c in range(c_an0, c_an1 + 1):
            L = get_column_letter(c)
            valeurs.append(f"=COUNTIFS(T_Programme[Code installation],$C{{r}},T_Programme[Type opération source],"
                           f"${C_TYPE}{{r}},T_Programme[Périodicité],$F{{r}},T_Programme[Année],{L}$7)")
        valeurs += [typ, f'=IFERROR(INDEX(T_Referentiel[Statut],{match}),"Absent du référentiel")',
                    f'=IF(COUNTIF({L_AN0}{{r}}:{L_AN1}{{r}},">1")>0,"Doublon potentiel : à confirmer","")']
        for c, v in enumerate(valeurs, start=1):
            if isinstance(v, str):
                v = v.replace("{r}", str(r))
            annee = c_an0 <= c <= c_an1
            gras = colonnes[c - 1] == "Opération"
            ecrire(ws, f"{get_column_letter(c)}{r}", v,
                   font=police(11 if annee else 10, gras or annee, VERT if gras or annee else TEXTE),
                   alignment=Alignment(horizontal="center" if annee or c == 6 else "left", vertical="center"),
                   number_format='[=1]"●";[>1]"●×"0;""' if annee else "General")
        ws.row_dimensions[r].height = 18
    ajouter_tableau(ws, t, f"A7:{der}{derniere}")

    an = f"{L_AN0}7"
    ws.conditional_formatting.add(f"{L_AN0}7:{L_AN1}7", regle(f"VALUE({an})=P_Annee", **JAUNE_TEXTE, stopIfTrue=True))
    ws.conditional_formatting.add(f"E{premiere}:{der}{derniere}", regle(
        f'OR(E{premiere}="À confirmer",E{premiere}="Absent du référentiel")', **JAUNE_TEXTE, stopIfTrue=True))
    ws.conditional_formatting.add(f"{C_VIG}{premiere}:{C_VIG}{derniere}", regle(
        f"LEN(${C_VIG}{premiere})>0", **JAUNE_TEXTE, stopIfTrue=True))
    ws.conditional_formatting.add(f"{L_AN0}{premiere}:{L_AN1}{derniere}", regle(
        f"VALUE({L_AN0}$7)=P_Annee", fill=fond(GRIS_MOYEN)))
    ws.conditional_formatting.add(f"A{premiere}:{der}{derniere}", regle(
        f"MOD(SUBTOTAL(103,$C${premiere}:$C{premiere}),2)=0", fill=fond(GRIS_CLAIR)))
    ws.conditional_formatting.add("C6", regle('ISNUMBER(SEARCH("ACTUALISER",$C$6))', fill=fond(JAUNE)))
    ws.freeze_panes = f"A{premiere}"
    mise_en_page(ws, f"A1:{der}{derniere}", "7:7")
    pied(ws)
    return len(cles), annees


# ---------------------------------------------------------------- 02_COUVERTURE_OT

def feuille_couverture(ws, colonnes, lignes):
    nb_aff = 11
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = JAUNE
    larg = [10, 10, 9, 13] + [10] * 5 + [14, 13] + [10] * 4
    largeurs(ws, larg)
    entete_vue(ws, nb_aff, "Couverture OT préventifs (GMAO)",
               "Nombre d'OT distincts par site, équipe, poste et activité. Complément du plan : "
               "aucun rapprochement automatique avec le programme.",
               etiquette=f"PROTOTYPE {VERSION}", mention="Données 100 % fictives")

    defs = [
        ("A", "B", '="OT distincts "&P_Annee',
         f"=IFERROR(SUMIFS({colonne_couv('OT distincts (toutes activités)')},{ANNEES_COUV},P_Annee),0)"),
        ("C", "D", '="Postes avec OT "&P_Annee', f"=COUNTIFS({ANNEES_COUV},P_Annee)"),
    ]
    for k in range(len(ACTIVITES)):
        col = get_column_letter(5 + k)
        libelle = f"INDEX(T_Ref_Activites[Activité],{k + 1})"
        defs.append((col, col, f"={libelle}",
                     f"=IFERROR(SUMIFS(INDEX({plage_couv()},0,MATCH({libelle},{ENTETE_COUV},0)),"
                     f"{ANNEES_COUV},P_Annee),0)"))
    defs += [
        ("J", "J", "OT toutes années", f"=IFERROR(SUM({colonne_couv('OT distincts (toutes activités)')}),0)"),
        ("K", "K", "Contrôles à vérifier", f'=COUNTIF({CTL_STATUTS},"À VÉRIFIER")'),
    ]
    tuiles(ws, 4, defs, groupes={"D", "I", "J"})
    ws["K4"].comment = Comment("Contrôles de l'import en statut « À VÉRIFIER » (détail : 08_PARAMETRES).",
                               "Prototype", width=240, height=60)
    ligne_info(ws, 6, [
        ("A6", "Année d'analyse", "etiquette"),
        ("B6", "=P_Annee", "valeur"),
        ("C6", f'="Source : "&{valeur_detail("Fichier source")}', "texte"),
        ("J6", "Actualisé le", "etiquette"),
        ("K6", f"={valeur_detail('Importé le')}", "valeur"),
    ])
    ws["B6"].number_format = "0"
    ws["K6"].number_format = "dd/mm/yyyy hh:mm"
    ecrire(ws, "A7", "Chaque chiffre = nombre d'OT distincts. Un OT à plusieurs activités compte une fois dans "
           "chacune de ses colonnes, et une seule fois dans « OT distincts (toutes activités) ».",
           font=police(8, couleur=TEXTE_DISCRET, italique=True), alignment=Alignment(vertical="center", indent=1))
    ws.row_dimensions[7].height = 18

    l0 = ANCRE_COUV[0]
    nb_col = len(colonnes)
    en_tete_tableau(ws, l0, colonnes, hauteur=34)
    for k, ligne in enumerate(lignes, start=l0 + 1):
        for c, nom in enumerate(colonnes, start=1):
            v = ligne[nom]
            num = c == 1 or c > 4
            ecrire(ws, f"{get_column_letter(c)}{k}", v,
                   font=police(10, c >= nb_col - 1, VERT if c >= nb_col - 1 else TEXTE),
                   alignment=Alignment(horizontal="center" if num else "left", vertical="center"),
                   number_format="0" if num else "General")
        ws.row_dimensions[k].height = 18
    fin = l0 + len(lignes)
    ws.auto_filter.ref = f"A{l0}:{get_column_letter(nb_col)}{fin}"

    zone = f"A{l0 + 1}:{MAX_COL}{l0 + 5000}"
    ws.conditional_formatting.add(f"{get_column_letter(nb_col)}{l0 + 1}:{MAX_COL}{l0 + 5000}", regle(
        f'AND(ISNUMBER({get_column_letter(nb_col)}{l0 + 1}),{get_column_letter(nb_col)}{l0 + 1}>0,'
        f'{get_column_letter(nb_col)}${l0}="OT avec point de contrôle")', **JAUNE_TEXTE, stopIfTrue=True))
    ws.conditional_formatting.add(f"E{l0 + 1}:{MAX_COL}{l0 + 5000}", regle(
        f'AND(ISNUMBER(E{l0 + 1}),E{l0 + 1}=0)', font=Font(name=POLICE, color=GRIS_MOYEN)))
    ws.conditional_formatting.add(zone, regle(f'AND($A{l0 + 1}<>"",$A{l0 + 1}=P_Annee)', fill=fond(GRIS_CLAIR)))
    ws.freeze_panes = f"A{l0 + 1}"
    mise_en_page(ws, f"A1:{get_column_letter(max(nb_col, nb_aff))}{fin}", f"{l0}:{l0}")
    pied(ws)


# ---------------------------------------------------------------- 03 / 04 : tables de détail

def feuille_programme(ws, cols, lignes):
    ws.sheet_properties.tabColor = VERT
    largeurs(ws, [12, 10, 8, 17, 18, 12, 22])
    en_tete_tableau(ws, 1, cols)
    for r, ligne in enumerate(lignes, start=2):
        for c, v in enumerate(ligne, start=1):
            ecrire(ws, f"{get_column_letter(c)}{r}", v, number_format="0" if isinstance(v, int) else "General")
    der = f"{get_column_letter(len(cols))}{len(lignes) + 1}"
    ajouter_tableau(ws, "T_Programme", f"A1:{der}")
    ws.freeze_panes = "A2"
    mise_en_page(ws, f"A1:{der}", "1:1", paysage=False)
    pied(ws)


def feuille_detail_ot(ws, detail):
    ws.sheet_properties.tabColor = VERT
    cols = gmao_ot.COLONNES_DETAIL
    larges = {"Contrôle qualité": 30, "Désignation": 44, "Désign.poste techn.": 22, "Écarts de structure": 24,
              "Fichier source": 36, "Poste technique": 17, "Importé le": 15, "Lignes export regroupées": 12}
    largeurs(ws, [larges.get(c, 11 if len(c) < 12 else 13) for c in cols])
    en_tete_tableau(ws, 1, cols, hauteur=34)
    for r, d in enumerate(detail, start=2):
        for c, nom in enumerate(cols, start=1):
            v = d[nom]
            fmt = ("dd/mm/yyyy hh:mm" if isinstance(v, dt.datetime) else "dd/mm/yyyy" if isinstance(v, dt.date)
                   else "#,##0.00" if nom == "Total coûts réels" else "0" if isinstance(v, int) else "General")
            ecrire(ws, f"{get_column_letter(c)}{r}", v, number_format=fmt)
    fin = len(detail) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{fin}"
    ws.freeze_panes = "C2"
    ws.conditional_formatting.add(f"G2:G{MAX_LIGNES}", regle("LEN($G2)>0", **JAUNE_TEXTE))
    mise_en_page(ws, f"A1:{get_column_letter(len(cols))}{fin}", "1:1")
    ws.page_setup.fitToWidth = 2
    pied(ws)


# ---------------------------------------------------------------- 05_REFERENTIEL

def feuille_referentiel(ws, ref_cols, ref_lignes, sites):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = VERT
    larg = [12, 20, 10, 12, 14, 30, 15, 3, 11, 13, 27, 26, 3, 10, 9, 27, 3, 16, 18, 14, 13, 24]
    largeurs(ws, larg)
    entete_vue(ws, len(larg), "Référentiels",
               "Correspondances de libellés. Corriger un référentiel corrige les vues ; rien n'est déduit des données.",
               etiquette=f"PROTOTYPE {VERSION}", mention="")
    blocs = [
        ("A", "Types d'opération du programme",
         "Lus par 01_PLAN_PREVENTIF (Opération, Statut). Copie du référentiel fictif."),
        ("I", "Activités GMAO",
         "Une activité absente d'ici est signalée « inconnue » à l'import. Libellés non fournis."),
        ("N", "Sites et équipes connus",
         "Tout autre couple est signalé « Nouveau site » ou « Nouvelle équipe »."),
        ("R", "Correspondance postes programme ↔ GMAO",
         "À construire avec le métier. Prépare la future comparaison prévu / OT ; non utilisée."),
    ]
    for col, titre, aide in blocs:
        ecrire(ws, f"{col}5", titre, font=police(11, True, VERT))
        ecrire(ws, f"{col}6", aide, font=police(8, couleur=TEXTE_DISCRET, italique=True),
               alignment=Alignment(vertical="top", wrap_text=False))
    ws.row_dimensions[6].height = 16
    lig = 8

    def table(col0, nom, entetes, lignes, largeurs_wrap):
        c0 = openpyxl.utils.column_index_from_string(col0)
        for k, h in enumerate(entetes):
            ecrire(ws, f"{get_column_letter(c0 + k)}{lig}", h, font=police(9.5, True, BLANC), fill=fond(VERT),
                   alignment=Alignment(vertical="center", wrap_text=True))
        for r, ligne in enumerate(lignes, start=lig + 1):
            for k, v in enumerate(ligne):
                ecrire(ws, f"{get_column_letter(c0 + k)}{r}", v, alignment=Alignment(vertical="center", wrap_text=True),
                       border=Border(bottom=FILET))
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 18,
                                              hauteur_texte(zip(ligne, largeurs_wrap)))
        fin = lig + max(1, len(lignes))
        ajouter_tableau(ws, nom, f"{col0}{lig}:{get_column_letter(c0 + len(entetes) - 1)}{fin}")
        return fin

    fins = [
        table("A", "T_Referentiel", ref_cols, ref_lignes, larg[0:7]),
        table("I", "T_Ref_Activites", ["Activité", "Libellé", "Statut", "Hypothèse de correspondance programme (non appliquée)"],
              [[a, "À confirmer", "Observée dans l'export modèle", HYPOTHESES_ACTIVITES.get(a, "")] for a in ACTIVITES],
              larg[8:12]),
        table("N", "T_Ref_Sites", ["Site", "Équipe", "Statut"],
              [[s, e, "Observé dans l'export modèle"] for s, e in sites], larg[13:16]),
        table("R", "T_Correspondance_Postes",
              ["Code installation (programme)", "Installation (programme)", "Code poste (GMAO)", "Statut", "Commentaire"],
              [[None, None, None, None, None]], larg[17:22]),
    ]
    ws.row_dimensions[lig].height = 34
    ws.conditional_formatting.add(f"A{lig + 1}:V{lig + 200}", regle(
        f'OR(A{lig + 1}="À confirmer",RIGHT(A{lig + 1},1)="?")', **JAUNE_TEXTE))
    ws.freeze_panes = f"A{lig + 1}"
    mise_en_page(ws, f"A1:V{max(fins)}", f"{lig}:{lig}")
    pied(ws)


# ---------------------------------------------------------------- 06_AVIS_UTILISATEURS

THEMES_AVIS = ["Utile", "Inutile", "Information manquante", "Filtre indispensable", "Besoin spécifique",
               "Saisie / commentaires", "Donnée GMAO à rechercher", "Couverture OT", "Autre"]


def feuille_avis(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = VERT
    larg = [12, 22, 14, 22, 50, 11, 22, 28]
    largeurs(ws, larg)
    entete_vue(ws, len(larg), "Avis utilisateurs", "Vos retours orientent la version suivante. Une ligne = une remarque.",
               mention="")
    ecrire(ws, "A5", "Questions pour guider vos retours", font=police(11, True, VERT))
    questions = [
        ("Utile", "Quelles colonnes ou informations vous servent vraiment ?"),
        ("Inutile", "Qu'est-ce qui pourrait être retiré ou simplifié ?"),
        ("Information manquante", "Que vous manque-t-il pour programmer l'année ?"),
        ("Filtre indispensable", "Sur quoi devez-vous pouvoir filtrer ou trier ?"),
        ("Besoin spécifique", "Un besoin propre à votre site ou à votre équipe ?"),
        ("Saisie / commentaires", "Avez-vous besoin de saisir ou de commenter ? Quoi, et qui ?"),
        ("Donnée GMAO à rechercher", "Quelles données faudrait-il récupérer dans l'extraction GMAO ?"),
        ("Couverture OT", "La vue 02 répond-elle à « mes OT préventifs sont-ils dans la GMAO ? » ?"),
    ]
    for k, (theme, question) in enumerate(questions, start=6):
        ecrire(ws, f"A{k}", "•", font=police(10, True, JAUNE), alignment=Alignment(horizontal="right"))
        ecrire(ws, f"B{k}", theme, font=police(9, True, VERT))
        ecrire(ws, f"D{k}", question, font=police(9))
    ecrire(ws, "A15", "Comment remplir : une ligne par remarque dans le tableau ci-dessous (cellules blanches). "
           "Thème, Priorité : listes déroulantes. La ligne grisée est un exemple à écraser.",
           font=police(8, couleur=TEXTE_DISCRET, italique=True))
    lig = 17
    colonnes = ["Date", "Contributeur (nom, fonction)", "Site / périmètre", "Thème", "Remarque",
                "Priorité", "Élément concerné (onglet, colonne)", "Suite donnée (équipe projet)"]
    en_tete_tableau(ws, lig, colonnes)
    exemple = [IMPORT_PROGRAMME_LE, "EXEMPLE – responsable d'exploitation", "PERIM_A", "Couverture OT",
               "Voir dans la couverture les postes du plan qui n'ont encore aucun OT.", "Moyenne",
               "02_COUVERTURE_OT", ""]
    nb_vides = 30
    for r in range(lig + 1, lig + 2 + nb_vides):
        ex = r == lig + 1
        for c in range(1, len(colonnes) + 1):
            ecrire(ws, f"{get_column_letter(c)}{r}", exemple[c - 1] if ex else None,
                   font=police(9, couleur=TEXTE_DISCRET, italique=True) if ex else police(10),
                   fill=fond(GRIS_CLAIR) if ex else fond(BLANC),
                   alignment=Alignment(vertical="top", wrap_text=True),
                   border=Border(bottom=FILET, left=FILET, right=FILET),
                   number_format="dd/mm/yyyy" if c == 1 else "General")
        ws.row_dimensions[r].height = 30
    fin = lig + 1 + nb_vides
    ajouter_tableau(ws, "T_Avis", f"A{lig}:H{fin}")
    for dv, col in (
        (DataValidation(type="list", formula1='"' + ",".join(THEMES_AVIS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Thème", error="Choisir un thème dans la liste."), "D"),
        (DataValidation(type="list", formula1='"Haute,Moyenne,Basse"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Priorité", error="Haute, Moyenne ou Basse."), "F"),
        (DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        showErrorMessage=True, errorTitle="Date", error="Saisir une date (jj/mm/aaaa)."), "A"),
    ):
        ws.add_data_validation(dv)
        dv.add(f"{col}{lig + 1}:{col}{fin}")
    mise_en_page(ws, f"A1:H{fin}", f"{lig}:{lig}")
    pied(ws)


# ---------------------------------------------------------------- 07_POINTS_A_VALIDER

def feuille_points(ws, points):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = VERT
    larg = [5, 15, 46, 34, 36, 13, 30]
    largeurs(ws, larg)
    entete_vue(ws, len(larg), "Points à valider",
               "Questions métier et hypothèses relevées dans les sources fictives. Rien n'a été corrigé ni supprimé.",
               mention="")
    ecrire(ws, "A5", f'="Contrôles automatiques de l\'import GMAO : "&COUNTIF({CTL_STATUTS},"À VÉRIFIER")'
                     f'&" à vérifier (détail dans 08_PARAMETRES)."',
           font=police(9, True, VERT), alignment=Alignment(vertical="center", indent=1))
    ws.conditional_formatting.add("A5", regle(f'COUNTIF({CTL_STATUTS},"À VÉRIFIER")>0', fill=fond(JAUNE)))
    lig = 7
    colonnes = ["N°", "Thème", "Constat (dans les sources)", "Traitement dans la V0.3", "Question à trancher",
                "Statut", "Décision / réponse"]
    en_tete_tableau(ws, lig, colonnes)
    for k, (theme, constat, traitement, question) in enumerate(points, start=1):
        r = lig + k
        ligne = [k, theme, constat, traitement, question, "À confirmer", None]
        for c, v in enumerate(ligne, start=1):
            ecrire(ws, f"{get_column_letter(c)}{r}", v,
                   font=police(9.5, c in (2, 6), VERT if c in (2, 6) else TEXTE),
                   alignment=Alignment(vertical="top", wrap_text=True, horizontal="center" if c == 1 else None),
                   border=Border(bottom=FILET))
        ws.row_dimensions[r].height = hauteur_texte(zip(ligne, larg))
    fin = lig + len(points)
    ajouter_tableau(ws, "T_Points", f"A{lig}:G{fin}")
    dv = DataValidation(type="list", formula1='"À confirmer,Validé,Rejeté,Reporté"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F{lig + 1}:F{fin + 30}")
    ws.freeze_panes = f"A{lig + 1}"
    mise_en_page(ws, f"A1:G{fin}", f"{lig}:{lig}")
    pied(ws)


# ---------------------------------------------------------------- 08_PARAMETRES

def feuille_parametres(wb, ws, controles):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GRIS_MOYEN
    larg = [5, 40, 12, 12, 12, 50, 14]
    largeurs(ws, larg)
    entete_vue(ws, len(larg), "Paramètres et contrôles de l'import",
               "Seules les cellules blanches encadrées se modifient. Power Query lit ces valeurs à chaque actualisation.",
               mention="")
    ecrire(ws, "B5", "Paramètres", font=police(11, True, VERT))
    dossier = 'IFERROR(LEFT(CELL("filename",$A$1),FIND("[",CELL("filename",$A$1))-1),"")'
    parametres = [
        # (ligne, libellé, valeur, nom défini, saisie ?, format, aide)
        (6, "Année d'analyse", ANNEE, "P_Annee", True, "0",
         "Année mise en avant dans 01 (colonne jaune) et 02 (indicateurs, lignes grisées)."),
        (7, "Dossier de ce classeur (automatique)", f"={dossier}", "P_DossierClasseur", False, "General",
         "Calculé par Excel quand le classeur est enregistré sur le poste (vide sinon)."),
        (8, "Nom du fichier d'extraction GMAO", FICHIER_GMAO, "P_NomFichierSource", True, "General",
         "Fichier .xlsx ou .csv (séparateur ;) placé dans le même dossier que ce classeur."),
        (9, "Chemin complet de l'extraction (facultatif)", None, "P_CheminComplet", True, "General",
         "Prioritaire s'il est rempli, ex. C:\\Maintenance\\Extraction_GMAO.xlsx. Pas d'adresse https."),
        (10, "Chemin utilisé par Power Query", '=IF(LEN(TRIM(P_CheminComplet))>0,TRIM(P_CheminComplet),'
         'P_DossierClasseur&P_NomFichierSource)', "P_CheminSource", False, "General", "Calculé : ne pas modifier."),
        (11, "Feuille de l'extraction (si fichier Excel)", None, "P_FeuilleSource", True, "General",
         "Vide = première feuille du fichier."),
        (12, "Longueur du code poste", 9, "P_LongueurCodePoste", True, "0",
         "Code poste = N premiers caractères du Poste technique (règle du test manuel : 9)."),
        (13, "Caractère attendu après le code poste", "-", "P_SeparateurCodePoste", True, "@",
         "Sert au contrôle de format (« - » observé dans l'export modèle). Vide = pas de contrôle."),
        (14, "Dernière actualisation des OT (automatique)", f"={valeur_detail('Importé le')}", None, False,
         "dd/mm/yyyy hh:mm", "Écrite par la requête OT_Detail à chaque « Actualiser tout »."),
        (15, "Source des OT actuellement chargée", f"={valeur_detail('Fichier source')}", None, False, "General",
         "« Instantané de démonstration » tant que les requêtes ne sont pas installées (09_NOTICE)."),
        (16, "Date d'import du programme préventif", IMPORT_PROGRAMME_LE, "P_DateProgramme", True, "dd/mm/yyyy",
         "Programme (03_DETAIL_PROGRAMME) : à mettre à jour à chaque nouvel import."),
    ]
    for ligne, libelle, valeur, nom, saisie, fmt, aide in parametres:
        ecrire(ws, f"B{ligne}", libelle, font=police(9.5, True, VERT), alignment=Alignment(vertical="center"))
        ws.merge_cells(f"C{ligne}:E{ligne}")
        ecrire(ws, f"C{ligne}", valeur, font=police(10, saisie, TEXTE if saisie else TEXTE_DISCRET),
               alignment=Alignment(horizontal="left", vertical="center", indent=1), number_format=fmt,
               fill=fond(BLANC) if saisie else fond(GRIS_CLAIR))
        for c in "CDE":
            ws[f"{c}{ligne}"].border = Border(top=FILET, bottom=FILET, left=FILET if c == "C" else None,
                                              right=FILET if c == "E" else None) if saisie else Border()
            if not saisie:
                ws[f"{c}{ligne}"].fill = fond(GRIS_CLAIR)
        ecrire(ws, f"F{ligne}", aide, font=police(8, couleur=TEXTE_DISCRET, italique=True),
               alignment=Alignment(vertical="center", wrap_text=True))
        ws.row_dimensions[ligne].height = 26
        if nom:
            wb.defined_names[nom] = DefinedName(nom, attr_text=f"'{S_PARAM}'!$C${ligne}")
    for dv, ref in (
        (DataValidation(type="whole", operator="between", formula1="2000", formula2="2100", showErrorMessage=True,
                        errorTitle="Année", error="Saisir une année sur 4 chiffres."), "C6"),
        (DataValidation(type="whole", operator="between", formula1="1", formula2="50", showErrorMessage=True,
                        errorTitle="Longueur", error="Nombre de caractères entre 1 et 50."), "C12"),
        (DataValidation(type="textLength", operator="lessThanOrEqual", formula1="1", showErrorMessage=True,
                        errorTitle="Séparateur", error="Un seul caractère, ou vide."), "C13"),
    ):
        ws.add_data_validation(dv)
        dv.add(ref)

    l0 = ANCRE_CTL[0]
    ecrire(ws, f"B{l0 - 2}", "Contrôles de l'import GMAO", font=police(11, True, VERT))
    ecrire(ws, f"B{l0 - 1}", "Une ligne par contrôle, recalculée à chaque actualisation. Les lignes concernées se "
           "filtrent dans 04_DETAIL_OT, colonne « Contrôle qualité ».",
           font=police(8, couleur=TEXTE_DISCRET, italique=True), alignment=Alignment(vertical="center"))
    en_tete_tableau(ws, l0, gmao_ot.COLONNES_CONTROLES, hauteur=30)
    for r, c in enumerate(controles, start=l0 + 1):
        for k, nom in enumerate(gmao_ot.COLONNES_CONTROLES, start=1):
            ecrire(ws, f"{get_column_letter(k)}{r}", c[nom],
                   font=police(9.5, nom in ("Contrôle", "Statut"), VERT if nom in ("Contrôle", "Statut") else TEXTE),
                   alignment=Alignment(vertical="center", wrap_text=nom == "Exemples (Ordre ou ligne)",
                                       horizontal="center" if k in (1, 4, 5) else None),
                   border=Border(bottom=FILET))
        ws.row_dimensions[r].height = 18
    fin = l0 + len(controles)
    ws.conditional_formatting.add(f"G{l0 + 1}:G{l0 + 60}", regle(f'G{l0 + 1}="À VÉRIFIER"', **JAUNE_TEXTE))
    ws.conditional_formatting.add(f"G{l0 + 1}:G{l0 + 60}", regle(f'G{l0 + 1}="INFO"', fill=fond(GRIS_CLAIR)))
    mise_en_page(ws, f"A1:G{fin}", paysage=False)
    pied(ws)


# ---------------------------------------------------------------- 09_NOTICE

def feuille_notice(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GRIS_MOYEN
    larg = [26, 62, 46]
    largeurs(ws, larg)
    entete_vue(ws, 3, "Notice", f"Prototype {VERSION} de programmation préventive et de couverture OT.", mention="")
    r = 4

    def section(titre):
        nonlocal r
        r += 1
        ecrire(ws, f"A{r}", titre, font=police(11, True, VERT))
        for c in range(1, 4):
            ws.cell(r, c).border = Border(bottom=Side(style="thin", color=JAUNE))
        ws.row_dimensions[r].height = 22
        r += 1

    def lignes(rangees, en_tete=None):
        nonlocal r
        if en_tete:
            for c, v in enumerate(en_tete, start=1):
                ecrire(ws, f"{get_column_letter(c)}{r}", v, font=police(8.5, True, TEXTE_DISCRET))
            r += 1
        for rangee in rangees:
            for c, v in enumerate(rangee, start=1):
                ecrire(ws, f"{get_column_letter(c)}{r}", v,
                       font=police(9.5, c == 1 and len(rangee) > 1, VERT if c == 1 and len(rangee) > 1 else TEXTE),
                       alignment=Alignment(vertical="top", wrap_text=True))
            if len(rangee) == 1:
                ws.merge_cells(f"A{r}:C{r}")
                ws.row_dimensions[r].height = hauteur_texte([(rangee[0], sum(larg))])
            else:
                ws.row_dimensions[r].height = hauteur_texte(zip(rangee, larg))
            r += 1

    section("En bref")
    lignes([
        ["Le produit est d'abord un outil de programmation et de suivi de la maintenance préventive par poste : "
         "Site → Équipe → Poste → Année → opérations prévues (01_PLAN_PREVENTIF)."],
        ["La couche OT GMAO (02_COUVERTURE_OT) répond en complément à : « les OT préventifs correspondant à ce que "
         "je dois faire sont-ils bien présents dans la GMAO ? ». Ce n'est pas un suivi de tous les OT : ni statuts, "
         "ni retards, ni clôtures."],
        ["Données 100 % fictives. La GMAO reste la source officielle : ne jamais corriger les données dans ce fichier."],
    ])
    section("Onglets")
    lignes([
        [S_PLAN, "Plan pluriannuel : une ligne par poste et opération, une colonne par année.", "Consultation."],
        [S_COUV, "Nombre d'OT distincts par année, site, équipe, poste et activité.", "Power Query (OT_Couverture)."],
        [S_PROG, "Programme préventif normalisé (T_Programme), une ligne par opération et par année.",
         "Remplacé à chaque import du programme."],
        [S_OT, "OT GMAO normalisés : une ligne par couple Ordre × Activité atomique.", "Power Query (OT_Detail)."],
        [S_REF, "Types du programme, activités GMAO, sites et équipes connus, correspondance des postes.",
         "Référent données."],
        [S_AVIS, "Recueil des retours, une ligne par remarque.", "Tous les testeurs."],
        [S_POINTS, "Questions métier à trancher.", "Équipe projet."],
        [S_PARAM, "Année, emplacement de l'extraction, règle du code poste, contrôles de l'import.",
         "Cellules blanches seulement."],
        [S_NOTICE, "Ce mode d'emploi.", ""],
        [S_CODE, "Code M des trois requêtes, à coller une seule fois (voir « Installer Power Query »).", ""],
    ], en_tete=["Onglet", "Contenu", "Mise à jour"])
    section("Règles de l'import GMAO")
    lignes([
        ["Site / Équipe", "Site = « Groupe gestionnaires » ; Équipe = « Poste travail princ. ».",
         "Correspondances observées lors du test manuel."],
        ["Code poste", "9 premiers caractères du « Poste technique » (Excel : =GAUCHE(Poste technique;9) ; "
         "Power Query : Text.Start([Poste technique], 9)). Longueur réglable dans 08.",
         "Poste vide, trop court ou de format inattendu : ligne gardée et signalée."],
        ["Année planifiée", "Année de la « Date début planifiée ».", "Date vide ou illisible : ligne gardée et signalée."],
        ["Activité", "Découpée sur « + », espaces retirés, en majuscules : « EC+PIS+FILTRE » donne 3 lignes "
         "(EC, PIS, FILTRE).", "Activité vide, mal formée ou inconnue du référentiel : signalée."],
        ["Doublons", "Un même couple Ordre + Activité n'est gardé qu'une fois ; la ligne gardée indique les lignes "
         "de l'export regroupées.", "Signalé : « Couple Ordre + Activité en double »."],
        ["Comptage", "Couverture = nombre d'OT DISTINCTS (numéro d'Ordre) par année, site, équipe, poste et activité.",
         "Un tableau croisé sur 04_DETAIL_OT peut utiliser « Nombre » : chaque couple n'y figure qu'une fois."],
        ["Colonnes", "Les 24 en-têtes de l'extraction sont repris tels quels (espaces de fin retirés, ex. « Typologie »).",
         "6 obligatoires : si l'une manque, l'actualisation s'arrête avec un message qui la nomme."],
    ], en_tete=["Règle", "Ce que fait l'import", "Cas particuliers"])
    section("Installer Power Query (une seule fois, Excel Desktop sous Windows)")
    lignes([
        ["1. Enregistrer ce classeur sur le poste, dans le dossier de l'extraction GMAO. Dans 08_PARAMETRES, "
         "vérifier le nom du fichier (ou saisir son chemin complet)."],
        ["2. Données › Obtenir des données › À partir d'autres sources › Requête vide. Accueil › Éditeur avancé : "
         "coller le code de OT_Detail (fichier OT_Detail.pq ou onglet 10_CODE_M), OK. Renommer la requête "
         "« OT_Detail » (volet Paramètres d'une requête › Nom)."],
        ["3. Recommencer avec OT_Couverture, puis OT_Controles (noms exacts : les requêtes s'appellent entre elles). "
         "Puis Accueil › Fermer et charger dans… › Créer uniquement la connexion."],
        ["4. Effacer les instantanés de démonstration : 04_DETAIL_OT, 02_COUVERTURE_OT à partir de A8 et "
         "08_PARAMETRES à partir de A20 (Données › Filtrer pour retirer les flèches, puis sélection et Suppr)."],
        ["5. Volet Requêtes et connexions : clic droit › Charger dans… › Tableau › Feuille de calcul existante : "
         "OT_Detail en ='04_DETAIL_OT'!$A$1, OT_Couverture en ='02_COUVERTURE_OT'!$A$8, "
         "OT_Controles en ='08_PARAMETRES'!$A$20."],
        ["Si Excel demande les niveaux de confidentialité, choisir « Organisationnel » pour le classeur et pour le "
         "fichier. Message « Formula.Firewall » : Données › Obtenir des données › Options de requête › Classeur actif "
         "› Confidentialité › Ignorer les niveaux de confidentialité (si la politique interne l'autorise)."],
    ])
    section("Actualiser (à chaque nouvelle extraction)")
    lignes([
        ["Remplacer le fichier d'extraction (même nom, ou nouveau nom dans 08_PARAMETRES), puis Données › "
         "Actualiser tout. Détail, couverture, contrôles et indicateurs se mettent à jour, y compris pour un nouvel "
         "OT, une nouvelle année, activité, site ou équipe."],
        ["Lire ensuite les contrôles (08_PARAMETRES) : tout statut « À VÉRIFIER » se filtre dans 04_DETAIL_OT, "
         "colonne « Contrôle qualité ». Ajouter au référentiel (05) un site, une équipe ou une activité validés."],
        ["Programme préventif : remplacer les lignes de 03_DETAIL_PROGRAMME, puis lire le contrôle de la ligne 6 de "
         "01_PLAN_PREVENTIF. En V0.3, une nouvelle combinaison poste / opération demande de régénérer la vue "
         "(outils/construire_prototype_v03.py)."],
    ])
    section("Power BI et comparaison future")
    lignes([
        ["Tables plates prêtes pour Power BI : T_Programme (03), le détail OT (04), T_Referentiel, T_Ref_Activites, "
         "T_Ref_Sites (05). La couverture (02) se recalcule dans Power BI par un nombre distinct d'Ordre."],
        ["Comparaison prévu / OT (prévu + OT présent, prévu sans OT, OT hors programme) : préparée mais pas activée. "
         "Il faut d'abord valider la clé de rapprochement (T_Correspondance_Postes et correspondance des activités, "
         "onglet 05)."],
    ])
    mise_en_page(ws, f"A1:C{r}", paysage=False)
    pied(ws)


# ---------------------------------------------------------------- 10_CODE_M

def feuille_code(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GRIS_MOYEN
    ws.column_dimensions["A"].width = 150
    entete_vue(ws, 1, "Code M des requêtes Power Query",
               "Copier un bloc (cellules de la colonne A, de « let » à la dernière ligne) dans l'Éditeur avancé.",
               etiquette="", mention="")
    r = 5
    for nom, cible in (("OT_Detail", "04_DETAIL_OT!A1"), ("OT_Couverture", "02_COUVERTURE_OT!A8"),
                       ("OT_Controles", "08_PARAMETRES!A20")):
        ecrire(ws, f"A{r}", f"Requête « {nom} » – fichier {nom}.pq – chargement : {cible}",
               font=police(11, True, VERT))
        ws[f"A{r}"].border = Border(bottom=Side(style="thin", color=JAUNE))
        r += 1
        for ligne in (CODE_M / f"{nom}.pq").read_text(encoding="utf-8").splitlines():
            c = ws.cell(r, 1, ligne if ligne.strip() else None)
            c.font = Font(name="Consolas", size=9, color=TEXTE)
            c.alignment = Alignment(vertical="center")
            ws.row_dimensions[r].height = 13
            r += 1
        r += 2
    pied(ws)


# ---------------------------------------------------------------- points à valider

def points_v03(lignes_prog, cols_prog, ref_lignes, nb_plan):
    i = {n: k for k, n in enumerate(cols_prog)}
    vus = {}
    for l in lignes_prog:
        k = tuple(l)
        vus[k] = vus.get(k, 0) + 1
    doublons = sorted((k for k, n in vus.items() if n > 1), key=lambda k: (k[i["Année"]], k[i["Code installation"]]))
    txt_doublons = " ; ".join(f"{k[i['Code installation']]} / {k[i['Type opération source']]} / {k[i['Année']]}"
                              for k in doublons)
    a_confirmer = [l for l in ref_lignes if "À confirmer" in [str(v) for v in l]]
    return [
        ("Hiérarchie",
         "Le programme fournit « Périmètre » et « Groupe » ; l'extraction GMAO, « Groupe gestionnaires » (site) et "
         "« Poste travail princ. » (équipe).",
         "Colonnes du programme affichées sous leurs noms d'origine.",
         "Périmètre = Site et Groupe = Équipe ? Avec quels libellés lisibles ?"),
        ("Rapprochement",
         "Les postes du programme (INST-xxx) et les codes poste GMAO (G-I-A1100…) n'ont aucune clé commune dans "
         "les sources fictives.",
         "Pas de comparaison prévu / OT. Table T_Correspondance_Postes préparée (05), vide.",
         "Quelle clé relie un poste du programme à un code poste GMAO ? Qui la maintient ?"),
        ("Activités",
         "Activités observées : " + ", ".join(ACTIVITES) + ". Aucun libellé ni correspondance avec les types du "
         "programme n'est fourni.",
         "Codes repris tels quels ; libellés « À confirmer » ; hypothèses FILTRE → TYPE_C et REV → TYPE_D notées, "
         "non appliquées.",
         "Que signifient EC, PIS, PILOTAGE ? Quelle correspondance avec les opérations du programme ?"),
        ("Code poste",
         "Règle du test manuel : 9 premiers caractères du Poste technique. Dans le modèle, le 10e caractère est "
         "toujours « - » (ex. G-I-A1100-FILT01).",
         "Règle paramétrée (08) ; poste vide, trop court ou sans « - » en 10e position signalé.",
         "La règle est-elle valable pour tous les sites ? Le séparateur est-il toujours « - » ?"),
        ("Site",
         "« Code site » (TECH-01…) correspond un pour un à « Groupe gestionnaires » (SITE01…) dans le modèle.",
         "Site = Groupe gestionnaires, comme dans le test manuel.",
         "Quel champ porte le site de référence ? « Code site » doit-il être affiché ?"),
        ("Équipe",
         "« Secteur exploitation » porte exactement les mêmes valeurs que « Poste travail princ. » sur les 90 lignes "
         "du modèle.",
         "Équipe = Poste travail princ., comme dans le test manuel.",
         "Les deux champs sont-ils toujours identiques dans l'extraction réelle ?"),
        ("Périmètre de l'extraction",
         "Toutes les lignes du modèle sont « Type d'ordre » ZPR2 et « Type de travail » PRE.",
         "Rien n'est filtré ; un Type de travail autre que PRE est signalé (Info).",
         "L'extraction réelle sera-t-elle limitée aux OT préventifs ? Sur quel critère ?"),
        ("Année planifiée",
         "L'année vient de « Date début planifiée » ; « Créé le » et « Date modif. fiche ordre » existent aussi.",
         "Année planifiée = année de la Date début planifiée.",
         "Est-ce la bonne date de référence pour rattacher un OT à une année du programme ?"),
        ("Statuts OT",
         "Statut utilisateur ACTI / TERM et Statut système présents dans l'extraction.",
         "Conservés dans 04 mais volontairement non exploités (hors périmètre du produit).",
         "Faut-il exclure certains statuts (ex. OT annulés) du comptage de couverture ?"),
        ("Référentiel programme",
         "; ".join(f"{l[0]} « {l[1]} » : référence technique et statut « À confirmer »" for l in a_confirmer) + ".",
         "Libellé affiché tel quel, statut « À confirmer » en jaune.",
         "Quel libellé, quelle référence technique pour ce type ?"),
        ("Doublons programme",
         f"Lignes strictement identiques dans le programme : {txt_doublons}.",
         "Conservées et signalées (●×2 dans le plan, « Doublon potentiel »).",
         "Deux interventions distinctes ou doublon d'extraction ?"),
        ("Périodicité",
         "Dans le programme, l'intervalle entre deux occurrences ne correspond pas toujours à la périodicité "
         "indiquée (ex. INST-017 Révision « 6 ans » en 2027 puis 2029).",
         "Aucune année n'est recalculée : seules les années de la source sont affichées.",
         "La périodicité est-elle indicative ? L'année prévue vient-elle uniquement de la GMAO ?"),
        ("Vue plan",
         f"La vue 01 compte {nb_plan} lignes poste × opération, construites à partir du programme.",
         "Les cases se recalculent ; une nouvelle combinaison poste / opération demande de régénérer la vue.",
         "Faut-il passer le programme en Power Query, comme les OT, une fois sa source réelle connue ?"),
    ]


# ---------------------------------------------------------------- main

def construire(sortie):
    cols_prog, lignes_prog = lire_table(SOURCES / "DATA_PDM_TEST.xlsx", "DONNEES")
    ref_cols, ref_lignes = lire_table(SOURCES / "REFERENTIEL_TEST.xlsx", "REFERENTIEL")
    modele = SOURCES / FICHIER_GMAO
    _, brut = gmao_ot.lire_export(modele)
    sites = sorted({(l["Groupe gestionnaires"].strip(), l["Poste travail princ."].strip()) for l in brut})
    detail = gmao_ot.ot_detail(modele, ACTIVITES, sites, 9, "-",
                               fichier_source=f"Instantané de démonstration – {FICHIER_GMAO}",
                               importe_le=INSTANTANE_LE)
    col_couv, couverture = gmao_ot.ot_couverture(detail)
    controles = gmao_ot.ot_controles(detail)

    wb = openpyxl.Workbook()
    wb._fonts = IndexedList([Font(name=POLICE, size=10, family=2)])
    wb.properties.title = "Prototype maintenance préventive V0.3"
    wb.properties.creator = f"Prototype {VERSION}"
    ws = wb.active
    ws.title = S_PLAN
    for nom in (S_COUV, S_PROG, S_OT, S_REF, S_AVIS, S_POINTS, S_PARAM, S_NOTICE, S_CODE):
        wb.create_sheet(nom)

    feuille_parametres(wb, wb[S_PARAM], controles)
    feuille_programme(wb[S_PROG], cols_prog, lignes_prog)
    feuille_detail_ot(wb[S_OT], detail)
    feuille_referentiel(wb[S_REF], ref_cols, ref_lignes, sites)
    nb_plan, annees = feuille_plan(ws, lignes_prog, cols_prog, ref_lignes, SOURCES / "logo_rvb.png")
    feuille_couverture(wb[S_COUV], col_couv, couverture)
    feuille_avis(wb[S_AVIS])
    feuille_points(wb[S_POINTS], points_v03(lignes_prog, cols_prog, ref_lignes, nb_plan))
    feuille_notice(wb[S_NOTICE])
    feuille_code(wb[S_CODE])
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    sortie.parent.mkdir(parents=True, exist_ok=True)
    wb.save(sortie)

    if shutil.which("soffice"):
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(sortie)],
                           check=True, capture_output=True, timeout=240)
            # le dossier du classeur dépend de la machine : il n'est pas figé dans le fichier livré
            figes = {(S_PARAM, "C7"): "", (S_PARAM, "C10"): FICHIER_GMAO}
            print(f"{injecter_valeurs(sortie, Path(tmp) / sortie.name, figes)} valeurs en cache injectées (LibreOffice)")
    else:
        print("LibreOffice absent : les formules seront calculées par Excel à l'ouverture.")
    print(f"Écrit : {sortie} — plan {nb_plan} lignes ({annees[0]}-{annees[-1]}), "
          f"{len(detail)} lignes OT × activité, {len(couverture)} lignes de couverture")
    # l'extraction modèle accompagne le classeur : le chemin par défaut (dossier du classeur) fonctionne
    shutil.copyfile(modele, sortie.parent / FICHIER_GMAO)


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--sortie", type=Path, default=SORTIE)
    construire(p.parse_args().sortie)
