"""Construit Prototype_Suivi_Maintenance_2027.xlsx à partir du pack de sources fictives.

Usage :
    python3 construire_prototype.py [--sources DOSSIER] [--sortie FICHIER] [--valeurs FICHIER_RECALCULE]

Le classeur produit ne contient que des données présentes dans les sources :
aucune valeur n'est inventée, corrigée ou supprimée. Les colonnes issues du
référentiel et les contrôles sont des formules, recalculées par Excel à l'ouverture.

Valeurs en cache : les résultats des formules sont recopiés dans le fichier pour que
les aperçus qui ne calculent pas (téléphone, messagerie, Power BI) affichent les bons
chiffres. Ils viennent d'une copie recalculée par LibreOffice (si `soffice` est
installé) ou du fichier passé avec --valeurs (copie recalculée par Excel, par exemple).
"""

import argparse
import datetime as dt
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from lxml import etree
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ICI = Path(__file__).resolve().parent
SOURCES = ICI.parent / "sources"
SORTIE = ICI.parent / "Prototype_Suivi_Maintenance_2027.xlsx"

ANNEE = 2027
DATE_ACTUALISATION = dt.datetime(2026, 10, 13)  # réception du jeu fictif DATA_PDM_TEST.xlsx

# Charte graphique simplifiée
POLICE = "Century Gothic"
JAUNE = "FFC600"
VERT = "214B4D"
BLANC = "FFFFFF"
GRIS_CLAIR = "E9EDED"
GRIS_MOYEN = "D3DBDB"
TEXTE = "2B2B2B"
TEXTE_DISCRET = "5F6B6B"

S_SUIVI = "01_SUIVI_2027"
S_SOURCE = "02_DONNEES_SOURCE"
S_REF = "03_REFERENTIEL"
S_AVIS = "04_AVIS_UTILISATEURS"
S_POINTS = "05_POINTS_A_VALIDER"
S_NOTICE = "06_NOTICE"


# ---------------------------------------------------------------- styles

def police(taille=10, gras=False, couleur=TEXTE, italique=False):
    return Font(name=POLICE, size=taille, bold=gras, italic=italique, color=couleur)


def fond(couleur):
    return PatternFill("solid", start_color=couleur, end_color=couleur)


FILET = Side(style="thin", color=GRIS_MOYEN)
SEPARATEUR_BLANC = Side(style="medium", color=BLANC)


def ecrire(ws, ref, valeur, **style):
    cell = ws[ref]
    cell.value = valeur
    cell.font = style.get("font", police())
    if "fill" in style:
        cell.fill = style["fill"]
    cell.alignment = style.get("alignment", Alignment(vertical="center"))
    if "border" in style:
        cell.border = style["border"]
    if "number_format" in style:
        cell.number_format = style["number_format"]
    return cell


def bandeau(ws, derniere_col, titre, sous_titre):
    """Bandeau vert canard (zone structurante) + filet jaune d'accent."""
    for col in range(1, derniere_col + 1):
        for row in (1, 2):
            ws.cell(row, col).fill = fond(VERT)
        ws.cell(3, col).fill = fond(JAUNE)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 3
    ecrire(ws, "A1", titre, font=police(16, True, BLANC), fill=fond(VERT),
           alignment=Alignment(vertical="bottom", indent=1))
    ecrire(ws, "A2", sous_titre, font=police(9, couleur=BLANC), fill=fond(VERT),
           alignment=Alignment(vertical="center", indent=1))


def en_tete_tableau(ws, row, colonnes, hauteur=30):
    for i, nom in enumerate(colonnes, start=1):
        ecrire(ws, f"{get_column_letter(i)}{row}", nom, font=police(9.5, True, BLANC),
               fill=fond(VERT), alignment=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[row].height = hauteur


def ajouter_tableau(ws, nom, ref):
    table = Table(displayName=nom, ref=ref)
    # Pas de style prédéfini Excel : la mise en forme suit la charte (cellules + MFC).
    table.tableStyleInfo = TableStyleInfo(name=None, showFirstColumn=False, showLastColumn=False,
                                          showRowStripes=False, showColumnStripes=False)
    ws.add_table(table)


def largeurs(ws, valeurs):
    for i, w in enumerate(valeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def hauteur_texte(textes_largeurs, ligne=13.5, mini=18):
    """Hauteur de ligne estimée pour des cellules à retour à la ligne."""
    lignes = 1
    for texte, largeur in textes_largeurs:
        if not texte:
            continue
        n = 0
        for paragraphe in str(texte).split("\n"):
            n += max(1, -(-len(paragraphe) // max(1, int(largeur * 0.95))))
        lignes = max(lignes, n)
    return max(mini, lignes * ligne + 6)


def mise_en_page(ws, zone, lignes_titres=None, paysage=True):
    ws.print_area = zone
    if lignes_titres:
        ws.print_title_rows = lignes_titres
    ws.page_setup.orientation = "landscape" if paysage else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.left.text = "Prototype V0.1 – données fictives – la GMAO reste la source de référence"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "&A – page &P / &N"
    ws.oddFooter.right.size = 8


# ---------------------------------------------------------------- lecture des sources

def lire_table(chemin, feuille):
    ws = openpyxl.load_workbook(chemin, data_only=True)[feuille]
    lignes = [list(r) for r in ws.iter_rows(values_only=True)]
    en_tetes = [str(v) for v in lignes[0]]
    donnees = [r for r in lignes[1:] if any(v is not None for v in r)]
    return en_tetes, donnees


# ---------------------------------------------------------------- feuilles

COLONNES_SUIVI = [
    # (en-tête, largeur, origine)
    ("Périmètre", 11, "source"),
    ("Groupe", 10, "source"),
    ("Code installation", 13, "source"),
    ("Installation", 18, "source"),
    ("Opération", 21, "referentiel"),
    ("Famille", 11, "referentiel"),
    ("Périodicité", 11, "source"),
    ("Référence technique", 13, "referentiel"),
    ("Code type source", 12, "source"),
    ("Statut référentiel", 17, "referentiel"),
    ("Point de vigilance", 29, "controle"),
]
LIGNE_EN_TETE_SUIVI = 7


def feuille_suivi(wb, ws, lignes_annee, nb_types, logo):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = JAUNE
    largeurs(ws, [c[1] for c in COLONNES_SUIVI])
    nb_col = len(COLONNES_SUIVI)
    der = get_column_letter(nb_col)

    # Bandeau : logo officiel (proportions d'origine), titre, mention GMAO
    bandeau(ws, nb_col, None, None)
    ws["A1"].value = ws["A2"].value = None
    ecrire(ws, "C1", '="Suivi annuel de maintenance "&AnneeSuivi', font=police(18, True, BLANC),
           fill=fond(VERT), alignment=Alignment(vertical="bottom"))
    ecrire(ws, "C2", "Vue de consultation des opérations prévues. La GMAO reste la source de "
           "référence : ce fichier n'en est qu'une lecture.", font=police(9, couleur=BLANC), fill=fond(VERT))
    ecrire(ws, f"{der}1", "PROTOTYPE V0.1", font=police(9, True, JAUNE), fill=fond(VERT),
           alignment=Alignment(horizontal="right", vertical="bottom"))
    ecrire(ws, f"{der}2", "Données 100 % fictives", font=police(9, couleur=BLANC, italique=True),
           fill=fond(VERT), alignment=Alignment(horizontal="right", vertical="center"))
    ajouter_logo(ws, logo, hauteur_px=60)

    # Indicateurs (lignes 4-5)
    premiere, derniere = LIGNE_EN_TETE_SUIVI + 1, LIGNE_EN_TETE_SUIVI + len(lignes_annee)
    t = "T_Suivi2027"
    tuiles = [
        ("A", "B", '="Opérations "&AnneeSuivi', f"=COUNTA({t}[Code installation])"),
        ("C", "C", "Installations",
         f'=SUMPRODUCT(({t}[Code installation]<>"")/COUNTIF({t}[Code installation],{t}[Code installation]&""))'),
        ("D", "D", "Affichées (filtres)", f"=SUBTOTAL(103,{t}[Code installation])"),
    ]
    for i in range(nb_types):
        col = get_column_letter(5 + i)
        lig = 2 + i
        tuiles.append((col, col, f"='{S_REF}'!$B${lig}", f"=COUNTIF({t}[Code type source],'{S_REF}'!$A${lig})"))
    tuiles += [
        ("J", "J", "Statut à confirmer",
         f'=COUNTIF({t}[Statut référentiel],"À confirmer")+COUNTIF({t}[Statut référentiel],"Absent du référentiel")'),
        ("K", "K", "Points de vigilance", f'=COUNTIF({t}[Point de vigilance],"?*")'),
    ]
    groupes = {"D", "I", "J"}  # séparation visuelle volume / répartition / qualité
    for debut, fin, libelle, formule in tuiles:
        for row, valeur, fnt, align in (
            (4, libelle, police(8, True, TEXTE_DISCRET), Alignment(horizontal="center", vertical="bottom", wrap_text=True)),
            (5, formule, police(16, True, VERT), Alignment(horizontal="center", vertical="center")),
        ):
            if debut != fin:
                ws.merge_cells(f"{debut}{row}:{fin}{row}")
            ecrire(ws, f"{debut}{row}", valeur, font=fnt, fill=fond(GRIS_CLAIR), alignment=align,
                   number_format="0")
            for c in range(openpyxl.utils.column_index_from_string(debut),
                           openpyxl.utils.column_index_from_string(fin) + 1):
                cell = ws.cell(row, c)
                cell.fill = fond(GRIS_CLAIR)
                droite = Side(style="thick", color=BLANC) if get_column_letter(c) in groupes else SEPARATEUR_BLANC
                cell.border = Border(left=SEPARATEUR_BLANC, right=droite)
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 26
    ws["D4"].comment = Comment("Nombre de lignes visibles avec les filtres en cours.", "Prototype")

    # Ligne d'information : année, actualisation, contrôle de cohérence, aide
    ecrire(ws, "A6", "Année suivie", font=police(8, couleur=TEXTE_DISCRET),
           alignment=Alignment(horizontal="right", vertical="center"))
    ecrire(ws, "B6", ANNEE, font=police(10, True, VERT), alignment=Alignment(horizontal="center", vertical="center"),
           number_format="0")
    ecrire(ws, "C6", "Actualisé le", font=police(8, couleur=TEXTE_DISCRET),
           alignment=Alignment(horizontal="right", vertical="center"))
    ecrire(ws, "D6", DATE_ACTUALISATION, font=police(10, True, VERT), number_format="dd/mm/yyyy",
           alignment=Alignment(horizontal="center", vertical="center"),
           border=Border(left=FILET, right=FILET, top=FILET, bottom=FILET))
    ws["D6"].comment = Comment(
        "Date du dernier chargement de 02_DONNEES_SOURCE.\n"
        "Ici : jeu fictif DATA_PDM_TEST.xlsx reçu le 13/10/2026.\n"
        "À mettre à jour à chaque import (date d'extraction GMAO).", "Prototype", width=260, height=90)
    ws.merge_cells("E6:I6")
    nb_src = f"COUNTIF(T_Source[Année],AnneeSuivi)"
    nb_vue = f"COUNTA({t}[Code installation])"
    ecrire(ws, "E6",
           f'=IF(AND({nb_src}={nb_vue},COUNTIF({t}[Point de vigilance],"Absente*")=0),'
           f'"Contrôle source : OK ("&{nb_vue}&" / "&{nb_src}&" opérations "&AnneeSuivi&")",'
           f'"Contrôle source : VUE À ACTUALISER (source "&{nb_src}&" / vue "&{nb_vue}&")")',
           font=police(8, True, VERT), alignment=Alignment(horizontal="left", vertical="center", indent=1))
    ws["E6"].comment = Comment(
        "Compare la vue 2027 à 02_DONNEES_SOURCE : même nombre de lignes pour l'année suivie "
        "et aucune ligne de la vue absente de la source.", "Prototype", width=260, height=80)
    ecrire(ws, f"{der}6", "Filtrer / rechercher : flèches d'en-tête  ·  jaune = à confirmer",
           font=police(8, couleur=TEXTE_DISCRET, italique=True),
           alignment=Alignment(horizontal="right", vertical="center"))
    ws.row_dimensions[6].height = 20

    # Tableau principal
    en_tete_tableau(ws, LIGNE_EN_TETE_SUIVI, [c[0] for c in COLONNES_SUIVI], hauteur=32)
    match = "MATCH($I{r},T_Referentiel[Valeur source],0)"
    for k, ligne in enumerate(lignes_annee):
        r = premiere + k
        valeurs = [
            ligne["Périmètre"], ligne["Groupe"], ligne["Code installation"], ligne["Installation"],
            f'=IFERROR(INDEX(T_Referentiel[Libellé affiché],{match}),"À confirmer")',
            f'=IFERROR(INDEX(T_Referentiel[Famille],{match}),"À confirmer")',
            ligne["Périodicité"],
            f'=IFERROR(INDEX(T_Referentiel[Référence technique fictive],{match}),"À confirmer")',
            ligne["Type opération source"],
            f'=IFERROR(INDEX(T_Referentiel[Statut],{match}),"Absent du référentiel")',
            f'=IF(COUNTIFS(T_Source[Code installation],$C{{r}},T_Source[Type opération source],$I{{r}},'
            f'T_Source[Périodicité],$G{{r}},T_Source[Année],AnneeSuivi)=0,"Absente de la source : actualiser",'
            f'IF(COUNTIFS({t}[Code installation],$C{{r}},{t}[Code type source],$I{{r}},{t}[Périodicité],$G{{r}})>1,'
            f'"Doublon potentiel : à confirmer",""))',
        ]
        for c, v in enumerate(valeurs, start=1):
            if isinstance(v, str):
                v = v.replace("{r}", str(r)).replace("{{r}}", str(r))
            gras = COLONNES_SUIVI[c - 1][0] == "Opération"
            centre = COLONNES_SUIVI[c - 1][0] in ("Périodicité",)
            ecrire(ws, f"{get_column_letter(c)}{r}", v, font=police(10, gras, VERT if gras else TEXTE),
                   alignment=Alignment(horizontal="center" if centre else "left", vertical="center"))
        ws.row_dimensions[r].height = 18
    ref = f"A{LIGNE_EN_TETE_SUIVI}:{der}{derniere}"
    ajouter_tableau(ws, t, ref)

    # Mises en forme conditionnelles : le texte porte toujours l'information, la couleur la souligne
    zone = f"A{premiere}:{der}{derniere}"
    jaune = dict(fill=fond(JAUNE), font=Font(name=POLICE, bold=True, color=VERT))
    ws.conditional_formatting.add(f"E{premiere}:J{derniere}", FormulaRule(
        formula=[f'OR(E{premiere}="À confirmer",E{premiere}="Absent du référentiel")'], stopIfTrue=True, **jaune))
    ws.conditional_formatting.add(f"K{premiere}:K{derniere}", FormulaRule(
        formula=[f"LEN($K{premiere})>0"], stopIfTrue=True, **jaune))
    ws.conditional_formatting.add(zone, FormulaRule(  # lignes alternées qui suivent les filtres
        formula=[f"MOD(SUBTOTAL(103,$C${premiere}:$C{premiere}),2)=0"], fill=fond(GRIS_CLAIR)))
    ws.conditional_formatting.add("E6", FormulaRule(
        formula=['ISNUMBER(SEARCH("ACTUALISER",$E$6))'], fill=fond(JAUNE)))

    ws.freeze_panes = f"A{premiere}"
    mise_en_page(ws, f"A1:{der}{derniere}", f"{LIGNE_EN_TETE_SUIVI}:{LIGNE_EN_TETE_SUIVI}")
    return ref


def ajouter_logo(ws, chemin, hauteur_px):
    """Logo inséré tel quel (fichier d'origine), largeur calculée sur le ratio exact."""
    img = Image(str(chemin))
    largeur_orig, hauteur_orig = img.width, img.height
    emu_h = hauteur_px * 9525
    emu_w = round(emu_h * largeur_orig / hauteur_orig)
    img.width, img.height = emu_w / 9525, hauteur_px
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=0, colOff=6 * 9525, row=0, rowOff=2 * 9525),
                               ext=XDRPositiveSize2D(emu_w, emu_h))
    ws.add_image(img)


def feuille_source(ws, en_tetes, lignes):
    ws.sheet_properties.tabColor = VERT
    largeurs(ws, [12, 10, 8, 17, 18, 12, 22])
    en_tete_tableau(ws, 1, en_tetes)
    for r, ligne in enumerate(lignes, start=2):
        for c, v in enumerate(ligne, start=1):
            ecrire(ws, f"{get_column_letter(c)}{r}", v, number_format="0" if isinstance(v, int) else "General")
    der = f"{get_column_letter(len(en_tetes))}{len(lignes) + 1}"
    ajouter_tableau(ws, "T_Source", f"A1:{der}")
    ws.freeze_panes = "A2"
    mise_en_page(ws, f"A1:{der}", "1:1", paysage=False)


def feuille_referentiel(ws, en_tetes, lignes):
    ws.sheet_properties.tabColor = VERT
    larg = [13, 21, 11, 13, 16, 44, 16]
    largeurs(ws, larg)
    en_tete_tableau(ws, 1, en_tetes)
    for r, ligne in enumerate(lignes, start=2):
        for c, v in enumerate(ligne, start=1):
            ecrire(ws, f"{get_column_letter(c)}{r}", v,
                   alignment=Alignment(vertical="center", wrap_text=True))
        ws.row_dimensions[r].height = hauteur_texte(zip(ligne, larg))
    der = f"{get_column_letter(len(en_tetes))}{len(lignes) + 1}"
    ajouter_tableau(ws, "T_Referentiel", f"A1:{der}")
    ws.conditional_formatting.add("A2:G200", FormulaRule(
        formula=['A2="À confirmer"'], fill=fond(JAUNE), font=Font(name=POLICE, bold=True, color=VERT)))
    ws.freeze_panes = "A2"
    ws.column_dimensions["H"].width = 3
    mise_en_page(ws, f"A1:{der}", "1:1")
    ws.column_dimensions["I"].width = 60
    notes = [
        ("Comment ce référentiel est utilisé", True),
        ("Les colonnes Opération, Famille, Référence technique et Statut de 01_SUIVI_2027 "
         "sont lues ici à partir du code « Valeur source ».", False),
        ("Corriger un libellé ici le corrige partout dans la vue.", False),
        ("Un code de la source absent de ce tableau apparaît « Absent du référentiel ».", False),
        ("« À confirmer » : information non validée, affichée telle quelle.", False),
    ]
    for i, (texte, titre) in enumerate(notes, start=1):
        ecrire(ws, f"I{i}", texte, font=police(9 if not titre else 10, titre, VERT if titre else TEXTE_DISCRET),
               alignment=Alignment(vertical="center", wrap_text=True))


THEMES_AVIS = ["Utile", "Inutile", "Information manquante", "Filtre indispensable", "Besoin spécifique",
               "Saisie / commentaires", "Donnée GMAO à rechercher", "Autre"]


def feuille_avis(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = VERT
    larg = [12, 22, 14, 22, 50, 11, 22, 28]
    largeurs(ws, larg)
    bandeau(ws, len(larg), "Avis utilisateurs", "Vos retours orientent la V0.2. Une ligne = une remarque.")

    ecrire(ws, "A5", "Questions pour guider vos retours", font=police(11, True, VERT))
    questions = [
        ("Utile", "Quelles colonnes ou informations vous servent vraiment ?"),
        ("Inutile", "Qu'est-ce qui pourrait être retiré ou simplifié ?"),
        ("Information manquante", "Que vous manque-t-il pour préparer l'année ?"),
        ("Filtre indispensable", "Sur quoi devez-vous pouvoir filtrer ou trier ?"),
        ("Besoin spécifique", "Un besoin propre à votre périmètre ou à votre métier ?"),
        ("Saisie / commentaires", "Avez-vous besoin de saisir ou de commenter ? Quoi, et qui ?"),
        ("Donnée GMAO à rechercher", "Quelles données faudrait-il récupérer dans l'extraction GMAO ?"),
    ]
    for i, (theme, question) in enumerate(questions, start=6):
        ecrire(ws, f"A{i}", "•", font=police(10, True, JAUNE), alignment=Alignment(horizontal="right"))
        ecrire(ws, f"B{i}", theme, font=police(9, True, VERT))
        ecrire(ws, f"D{i}", question, font=police(9))
    ecrire(ws, "A14", "Comment remplir : une ligne par remarque dans le tableau ci-dessous (cellules blanches). "
           "Thème, Priorité : listes déroulantes. La ligne grisée est un exemple à écraser.",
           font=police(8, couleur=TEXTE_DISCRET, italique=True))

    lig = 16
    colonnes = ["Date", "Contributeur (nom, fonction)", "Périmètre / site", "Thème", "Remarque",
                "Priorité", "Élément concerné (onglet, colonne)", "Suite donnée (équipe projet)"]
    en_tete_tableau(ws, lig, colonnes)
    exemple = [DATE_ACTUALISATION, "EXEMPLE – responsable d'exploitation", "PERIM_A", "Filtre indispensable",
               "Pouvoir filtrer par tournée pour préparer les passages sur le terrain.", "Moyenne",
               "01_SUIVI_2027 – filtres", ""]
    nb_vides = 30
    for r in range(lig + 1, lig + 2 + nb_vides):
        for c in range(1, len(colonnes) + 1):
            v = exemple[c - 1] if r == lig + 1 else None
            ecrire(ws, f"{get_column_letter(c)}{r}", v,
                   font=police(9, couleur=TEXTE_DISCRET, italique=True) if r == lig + 1 else police(10),
                   fill=fond(GRIS_CLAIR) if r == lig + 1 else fond(BLANC),
                   alignment=Alignment(vertical="top", wrap_text=True),
                   border=Border(bottom=FILET, left=FILET, right=FILET),
                   number_format="dd/mm/yyyy" if c == 1 else "General")
        ws.row_dimensions[r].height = 30
    fin = lig + 1 + nb_vides
    ajouter_tableau(ws, "T_Avis", f"A{lig}:{get_column_letter(len(colonnes))}{fin}")

    dv_theme = DataValidation(type="list", formula1='"' + ",".join(THEMES_AVIS) + '"', allow_blank=True,
                              showErrorMessage=True, errorTitle="Thème",
                              error="Choisir un thème dans la liste.",
                              promptTitle="Thème", prompt="Utile, inutile, manquant, filtre, besoin, saisie, donnée GMAO…",
                              showInputMessage=True)
    dv_prio = DataValidation(type="list", formula1='"Haute,Moyenne,Basse"', allow_blank=True,
                             showErrorMessage=True, errorTitle="Priorité", error="Haute, Moyenne ou Basse.")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                             showErrorMessage=True, errorTitle="Date", error="Saisir une date (jj/mm/aaaa).")
    for dv, col in ((dv_theme, "D"), (dv_prio, "F"), (dv_date, "A")):
        ws.add_data_validation(dv)
        dv.add(f"{col}{lig + 1}:{col}{fin}")
    mise_en_page(ws, f"A1:H{fin}", f"{lig}:{lig}")


def feuille_points(ws, points):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = VERT
    larg = [5, 14, 46, 34, 34, 13, 30]
    largeurs(ws, larg)
    bandeau(ws, len(larg), "Points à valider",
            "Hypothèses et incohérences relevées dans les sources fictives. Rien n'a été corrigé ni supprimé.")
    lig = 5
    colonnes = ["N°", "Thème", "Constat (dans les sources)", "Traitement dans la V0.1", "Question à trancher",
                "Statut", "Décision / réponse"]
    en_tete_tableau(ws, lig, colonnes)
    for k, (theme, constat, traitement, question) in enumerate(points, start=1):
        r = lig + k
        ligne = [k, theme, constat, traitement, question, "À confirmer", None]
        for c, v in enumerate(ligne, start=1):
            ecrire(ws, f"{get_column_letter(c)}{r}", v,
                   font=police(9.5, c in (2, 6), VERT if c in (2, 6) else TEXTE),
                   alignment=Alignment(vertical="top", wrap_text=True, horizontal="center" if c == 1 else None),
                   border=Border(bottom=FILET))
        ws.row_dimensions[r].height = hauteur_texte(zip(ligne, larg))
    fin = lig + len(points)
    ajouter_tableau(ws, "T_Points", f"A{lig}:G{fin}")
    dv = DataValidation(type="list", formula1='"À confirmer,Validé,Rejeté,Reporté V0.2"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F{lig + 1}:F{fin + 20}")
    ws.freeze_panes = f"A{lig + 1}"
    mise_en_page(ws, f"A1:G{fin}", f"{lig}:{lig}")


def feuille_notice(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GRIS_MOYEN
    larg = [24, 62, 44]
    largeurs(ws, larg)
    bandeau(ws, len(larg), "Notice", "Prototype V0.1 de suivi annuel de maintenance – à lire avant diffusion.")
    r = 5

    def section(titre):
        nonlocal r
        r += 1
        ecrire(ws, f"A{r}", titre, font=police(11, True, VERT))
        for c in range(1, 4):
            ws.cell(r, c).border = Border(bottom=Side(style="thin", color=JAUNE))
        ws.row_dimensions[r].height = 22
        r += 1

    def lignes(rangees, en_tete=None):
        nonlocal r
        if en_tete:
            for c, v in enumerate(en_tete, start=1):
                ecrire(ws, f"{get_column_letter(c)}{r}", v, font=police(8.5, True, TEXTE_DISCRET))
            r += 1
        for rangee in rangees:
            for c, v in enumerate(rangee, start=1):
                ecrire(ws, f"{get_column_letter(c)}{r}", v,
                       font=police(9.5, c == 1 and len(rangee) > 1, VERT if c == 1 and len(rangee) > 1 else TEXTE),
                       alignment=Alignment(vertical="top", wrap_text=True))
            if len(rangee) == 1:
                ws.merge_cells(f"A{r}:C{r}")
                ws.row_dimensions[r].height = hauteur_texte([(rangee[0], sum(larg))])
            else:
                ws.row_dimensions[r].height = hauteur_texte(zip(rangee, larg))
            r += 1

    section("En bref")
    lignes([
        ["Ce classeur répond à une question : « Qu'est-ce que j'ai à faire en maintenance en 2027 ? ». "
         "Il sert à recueillir l'avis des responsables avant toute industrialisation."],
        ["Les données sont 100 % fictives. La future GMAO restera la source officielle : ce fichier n'est qu'une "
         "lecture, il ne doit pas devenir une seconde source de vérité (ne pas y corriger les données)."],
    ])
    section("Onglets")
    lignes([
        [S_SUIVI, "Vue de consultation des opérations de l'année suivie, enrichie par le référentiel.",
         "Personne : consultation, filtres et tris."],
        [S_SOURCE, "Copie tabulaire de l'export (toutes années). Une ligne = une opération sur une installation.",
         "Remplacé à chaque import. Jamais modifié à la main."],
        [S_REF, "Correspondance code source → libellé, famille, référence technique, statut.",
         "Référent données."],
        [S_AVIS, "Recueil des retours, une ligne par remarque.", "Tous les testeurs."],
        [S_POINTS, "Hypothèses et incohérences à trancher avant la V0.2.", "Équipe projet."],
    ], en_tete=["Onglet", "Rôle", "Qui le modifie"])
    section("Colonnes du tableau principal")
    lignes([
        ["Périmètre, Groupe", "Source (02) – codes tels qu'exportés.", "Pas de libellé lisible fourni : à confirmer."],
        ["Code installation", "Source (02) – identifiant de l'installation.", ""],
        ["Installation", "Source (02) – libellé de l'installation.", ""],
        ["Opération", "Référentiel (03) – « Libellé affiché » du code type.", "« À confirmer » si le code est inconnu."],
        ["Famille", "Référentiel (03).", ""],
        ["Périodicité", "Source (02) – texte tel qu'exporté.",
         "Indicative : aucune année n'est recalculée à partir d'elle."],
        ["Référence technique", "Référentiel (03) – rattachée au type d'opération.", ""],
        ["Code type source", "Source (02) – « Type opération source ».", "Garde le lien avec la GMAO."],
        ["Statut référentiel", "Référentiel (03) – « Statut ».", "« Absent du référentiel » si le code est inconnu."],
        ["Point de vigilance", "Contrôle calculé.",
         "« Doublon potentiel » : même installation, même type, même périodicité, même année. "
         "« Absente de la source » : la vue n'est plus à jour."],
    ], en_tete=["Colonne", "Origine", "Remarque"])
    section("Lire le tableau")
    lignes([
        ["Filtres et recherche : flèches des en-têtes (zone de recherche intégrée), ou Ctrl + F."],
        ["« Affichées (filtres) » compte les lignes visibles ; les autres indicateurs portent sur toute l'année."],
        ["Jaune = information à confirmer. La couleur ne fait que souligner : le texte de la cellule le dit toujours."],
        ["Seules cellules à modifier dans 01_SUIVI_2027 : la date « Actualisé le » (D6). Rien d'autre n'est à saisir."],
    ])
    section("Actualiser les données")
    lignes([
        ["1. Remplacer les lignes de 02_DONNEES_SOURCE par le nouvel export (mêmes colonnes, à partir de A2)."],
        ["2. Mettre à jour la date « Actualisé le » (01_SUIVI_2027, cellule D6)."],
        ["3. Lire le contrôle source (01_SUIVI_2027, ligne 6). « VUE À ACTUALISER » signifie que la vue 2027 ne "
         "correspond plus à la source : en V0.1, elle est régénérée par le concepteur du prototype "
         "(outils/construire_prototype.py)."],
        ["Cible V0.2 : une requête Power Query (export GMAO → filtre sur l'année → fusion avec le référentiel) "
         "remplace l'étape 3 ; l'actualisation se fait alors par Données › Actualiser tout."],
    ])
    section("Compatibilité Power BI")
    lignes([
        ["T_Source (02) et T_Referentiel (03) sont des tableaux plats sans cellule fusionnée ni formule. "
         "Relation : T_Source[Type opération source] → T_Referentiel[Valeur source]. "
         "T_Avis (04) et T_Points (05) sont aussi importables."],
    ])
    mise_en_page(ws, f"A1:C{r}", paysage=False)


# ---------------------------------------------------------------- valeurs en cache

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def injecter_valeurs(classeur, recalcule, remplacements=None):
    """Recopie les résultats d'une copie recalculée comme valeurs en cache des formules.

    remplacements : {(feuille, cellule): valeur} imposée en cache, par exemple pour ne pas
    figer un chemin propre à la machine de construction.
    """
    remplacements = remplacements or {}
    valeurs = openpyxl.load_workbook(recalcule, data_only=True)
    with zipfile.ZipFile(classeur) as z:
        contenu = {n: z.read(n) for n in z.namelist()}
    wbxml = etree.fromstring(contenu["xl/workbook.xml"])
    rels = etree.fromstring(contenu["xl/_rels/workbook.xml.rels"])
    cibles = {r.get("Id"): r.get("Target") for r in rels}
    nb = 0
    for sheet in wbxml.iter(f"{{{NS}}}sheet"):
        cible = cibles[sheet.get(f"{{{NS_R}}}id")].lstrip("/")
        chemin = cible if cible.startswith("xl/") else f"xl/{cible}"
        ws_val = valeurs[sheet.get("name")]
        xml = etree.fromstring(contenu[chemin])
        for c in xml.iter(f"{{{NS}}}c"):
            f = c.find(f"{{{NS}}}f")
            if f is None:
                continue
            for ancien in c.findall(f"{{{NS}}}v"):
                c.remove(ancien)
            val = remplacements.get((sheet.get("name"), c.get("r")), ws_val[c.get("r")].value)
            if isinstance(val, dt.datetime):
                val = (val - dt.datetime(1899, 12, 30)).total_seconds() / 86400
            elif isinstance(val, dt.date):
                val = (val - dt.date(1899, 12, 30)).days
            v = etree.SubElement(c, f"{{{NS}}}v")
            if isinstance(val, bool):
                c.set("t", "b")
                v.text = "1" if val else "0"
            elif isinstance(val, (int, float)):
                c.attrib.pop("t", None)
                v.text = repr(float(val)) if isinstance(val, float) and not val.is_integer() else str(int(val))
            else:
                c.set("t", "str")
                v.text = "" if val is None else str(val)
            nb += 1
        contenu[chemin] = etree.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True)
    tmp = Path(str(classeur) + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in contenu.items():
            z.writestr(n, data)
    shutil.move(tmp, classeur)
    return nb


# ---------------------------------------------------------------- points à valider

def points_a_valider(lignes_source, cols, lignes_annee, ref_lignes, ref_cols):
    i = {nom: k for k, nom in enumerate(cols)}
    cle = lambda l: (l[i["Code installation"]], l[i["Type opération source"]], l[i["Périodicité"]], l[i["Année"]])
    vus, doublons = {}, []
    for l in lignes_source:
        vus[cle(l)] = vus.get(cle(l), 0) + 1
    doublons = sorted((k for k, n in vus.items() if n > 1), key=lambda k: (k[3], k[0]))
    libelle = {l[ref_cols.index("Valeur source")]: l[ref_cols.index("Libellé affiché")] for l in ref_lignes}
    a_confirmer = [l for l in ref_lignes if "À confirmer" in [str(v) for v in l]]
    nb_annee_doublons = sum(vus[k] for k in doublons if k[3] == ANNEE)

    # Intervalles entre occurrences successives différents de la périodicité affichée
    occ = {}
    for l in lignes_source:
        occ.setdefault((l[i["Code installation"]], l[i["Type opération source"]], l[i["Périodicité"]]), set()).add(
            l[i["Année"]])
    ecarts = []
    for (code, typ, perio), annees in sorted(occ.items()):
        a = sorted(annees)
        n = int(re.match(r"\d+", perio).group())
        for x, y in zip(a, a[1:]):
            if y - x != n:
                ecarts.append(f"{code} {libelle.get(typ, typ)} « {perio} » : {x} puis {y}")
    annees = sorted({l[i["Année"]] for l in lignes_source})

    txt_doublons = " ; ".join(f"{c} / {t} / {a} ({vus[(c, t, p, a)]} lignes)" for c, t, p, a in doublons)
    return [
        ("Référentiel",
         "; ".join(f"{l[0]} « {l[1]} » : " + ", ".join(
             f"{ref_cols[k].lower().replace(' fictive', '')} « À confirmer »" for k, v in enumerate(l) if v == "À confirmer")
                   for l in a_confirmer)
         + f". {sum(1 for l in lignes_annee if l['Type opération source'] in [x[0] for x in a_confirmer])} "
           f"opérations {ANNEE} concernées.",
         "Libellé affiché tel quel ; référence et statut « À confirmer » (cellules jaunes).",
         "Quel libellé, quelle famille et quelle référence technique ? Ce type existe-t-il dans la GMAO ?"),
        ("Doublons",
         f"Lignes strictement identiques dans la source : {txt_doublons}.",
         f"Toutes les lignes sont conservées ; les {nb_annee_doublons} lignes {ANNEE} concernées sont signalées "
         "« Doublon potentiel ».",
         "Deux interventions distinctes ou doublon d'extraction ? Quel identifiant unique (n° de gamme, d'OT) "
         "l'export GMAO fournira-t-il ?"),
        ("Périodicité",
         "L'intervalle entre deux occurrences ne correspond pas toujours à la périodicité indiquée, ex. "
         + " ; ".join(ecarts[:4]) + f" ({len(ecarts)} cas).",
         "Aucune année n'est recalculée : seules les années présentes dans la source sont affichées.",
         "La périodicité est-elle indicative ? L'année prévue doit-elle venir uniquement de la GMAO ?"),
        ("Ancien suivi",
         "Les mêmes codes portent des libellés différents dans l'ancien suivi et dans la source "
         "(ex. INST-001 = « Installation Démo 01 » / « Installation A1 »).",
         "Aucun rapprochement entre l'ancien suivi et la source.",
         "Le code installation est-il l'identifiant commun fiable entre les outils ?"),
        ("Ancien suivi",
         "Les colonnes « contrôle », « révision » et « contrôle filtre » de l'ancien suivi semblent correspondre à "
         "TYPE_A, TYPE_D et TYPE_C, mais aucune table de correspondance ne le confirme.",
         "Hypothèse non appliquée.",
         "Confirmer la correspondance avant toute reprise de l'historique."),
        ("Référentiel",
         "Périmètre et Groupe ne sont fournis que sous forme de codes (PERIM_A, GRP_A1…).",
         "Codes affichés tels quels.",
         "Faut-il un référentiel des périmètres et groupes (libellés lisibles, responsables) ?"),
        ("Référentiel",
         "La référence technique est portée par le type d'opération, pas par l'installation.",
         "Même référence pour toutes les opérations d'un même type.",
         "Dans la réalité, la référence dépend-elle de l'installation ou du matériel ?"),
        ("Granularité",
         "La source ne donne que l'année ; l'ancien suivi donnait des dates (dernier / prochain passage).",
         "Vue à l'année, sans mois ni date.",
         "Un mois ou une date prévisionnelle est-il nécessaire ? Est-il disponible dans la GMAO ?"),
        ("Périmètre de la vue",
         f"La source couvre {annees[0]} à {annees[-1]} ({len(lignes_source)} lignes) ; "
         f"{len(lignes_annee)} lignes concernent {ANNEE}.",
         f"Toutes les années restent dans 02_DONNEES_SOURCE ; la vue n'affiche que {ANNEE}.",
         "Une vue N+1 ou pluriannuelle serait-elle utile ?"),
    ]


# ---------------------------------------------------------------- main

def construire(sources, sortie, recalcule=None):
    cols, lignes_source = lire_table(sources / "DATA_PDM_TEST.xlsx", "DONNEES")
    ref_cols, ref_lignes = lire_table(sources / "REFERENTIEL_TEST.xlsx", "REFERENTIEL")
    i = {nom: k for k, nom in enumerate(cols)}
    ordre_type = {l[0]: k for k, l in enumerate(ref_lignes)}
    lignes_annee = [dict(zip(cols, l)) for l in lignes_source if l[i["Année"]] == ANNEE]
    lignes_annee.sort(key=lambda d: (d["Périmètre"], d["Groupe"], d["Code installation"],
                                     ordre_type.get(d["Type opération source"], 99)))

    wb = openpyxl.Workbook()
    wb._fonts = IndexedList([Font(name=POLICE, size=10, family=2)])  # police par défaut du classeur
    wb.properties.title = "Prototype suivi maintenance 2027"
    wb.properties.creator = "Prototype V0.1"
    ws = wb.active
    ws.title = S_SUIVI
    for nom in (S_SOURCE, S_REF, S_AVIS, S_POINTS, S_NOTICE):
        wb.create_sheet(nom)
    wb.defined_names["AnneeSuivi"] = DefinedName("AnneeSuivi", attr_text=f"'{S_SUIVI}'!$B$6")

    feuille_source(wb[S_SOURCE], cols, lignes_source)
    feuille_referentiel(wb[S_REF], ref_cols, ref_lignes)
    feuille_suivi(wb, ws, lignes_annee, len(ref_lignes), sources / "logo_rvb.png")
    feuille_avis(wb[S_AVIS])
    feuille_points(wb[S_POINTS], points_a_valider(lignes_source, cols, lignes_annee, ref_lignes, ref_cols))
    feuille_notice(wb[S_NOTICE])
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(sortie)
    if recalcule:
        print(f"{injecter_valeurs(sortie, recalcule)} valeurs en cache injectées")
    elif shutil.which("soffice"):
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(sortie)],
                           check=True, capture_output=True, timeout=180)
            print(f"{injecter_valeurs(sortie, Path(tmp) / sortie.name)} valeurs en cache injectées (LibreOffice)")
    else:
        print("LibreOffice absent : les formules seront calculées par Excel à l'ouverture.")
    print(f"Écrit : {sortie} ({len(lignes_annee)} opérations {ANNEE}, {len(lignes_source)} lignes source)")


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--sources", type=Path, default=SOURCES)
    p.add_argument("--sortie", type=Path, default=SORTIE)
    p.add_argument("--valeurs", type=Path, help="copie recalculée dont les résultats sont mis en cache")
    a = p.parse_args()
    construire(a.sources, a.sortie, a.valeurs)
