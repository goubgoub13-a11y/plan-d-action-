"""Transposition Python des requêtes Power Query de la V0.3 (dossier V03/power-query).

Sert à deux choses :
- produire l'instantané de démonstration chargé dans le classeur (mêmes colonnes,
  même ordre, mêmes règles que les requêtes M) ;
- tester les règles (éclatement des activités, comptage distinct, contrôles) sur
  des fichiers d'essai, faute de moteur Power Query dans l'environnement de build.

La référence fonctionnelle reste le code M. Toute évolution d'une règle doit être
faite dans les deux.
"""

import csv
import datetime as dt
import re
from pathlib import Path

import openpyxl

OBLIGATOIRES = ["Ordre", "Poste technique", "Date début planifiée", "Activité",
                "Groupe gestionnaires", "Poste travail princ."]
ATTENDUES = ["Créé le", "Ordre", "Ordre supérieur", "Type d'ordre", "Type de travail", "Avis",
             "Désignation", "Statut utilisateur", "Statut système", "Saisi par", "Total coûts réels",
             "Poste travail princ.", "Groupe gestionnaires", "Secteur exploitation", "Date début planifiée",
             "Code site", "Poste technique", "Désign.poste techn.", "Équipement", "Groupe de gamme principale",
             "Compteur de groupe de gamme principale", "Date modif. fiche ordre", "Typologie", "Activité"]
COLONNES_DATES = ["Créé le", "Date début planifiée", "Date modif. fiche ordre"]
COLONNES_NOMBRES = ["Total coûts réels"]

COLONNES_DETAIL = (
    ["Ordre", "Activité atomique", "Année planifiée", "Groupe gestionnaires", "Poste travail princ.",
     "Code poste", "Contrôle qualité", "Ligne export"]
    + [c for c in ATTENDUES if c not in ("Ordre", "Groupe gestionnaires", "Poste travail princ.")]
    + ["Ordre d'affichage activité", "Nb lignes export de l'OT", "Occurrences Ordre+Activité",
       "Lignes export regroupées", "Écarts de structure", "Fichier source", "Importé le"]
)

CONTROLES = [  # (libellé dans « Contrôle qualité », niveau)
    ("Ordre vide", "À vérifier"),
    ("OT présent sur plusieurs lignes de l'export", "À vérifier"),
    ("Poste technique vide", "À vérifier"),
    ("Poste technique trop court", "À vérifier"),
    ("Format du Poste technique inattendu", "À vérifier"),
    ("Date début planifiée vide", "À vérifier"),
    ("Date début planifiée illisible", "À vérifier"),
    ("Activité vide", "À vérifier"),
    ("Activité mal formée", "À vérifier"),
    ("Activité inconnue du référentiel", "À vérifier"),
    ("Couple Ordre + Activité en double", "Info"),
    ("Site ou équipe vide", "À vérifier"),
    ("Nouveau site", "À vérifier"),
    ("Nouvelle équipe", "À vérifier"),
    ("Type de travail autre que PRE", "Info"),
]
COLONNES_CONTROLES = ["N°", "Contrôle", "Niveau", "Lignes de l'export", "OT distincts",
                      "Exemples (Ordre ou ligne)", "Statut"]


class ErreurImport(Exception):
    """Équivalent des erreurs explicites levées par la requête OT_Detail."""


# ---------------------------------------------------------------- lecture

def lire_export(chemin, feuille=""):
    chemin = Path(chemin)
    if chemin.suffix.lower() in (".csv", ".txt"):
        with open(chemin, encoding="utf-8-sig", newline="") as f:
            lignes = list(csv.reader(f, delimiter=";"))
    else:
        wb = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
        if feuille and feuille not in wb.sheetnames:
            raise ErreurImport(f"Aucune feuille « {feuille} » dans l'extraction GMAO.")
        ws = wb[feuille] if feuille else wb.worksheets[0]
        lignes = [list(r) for r in ws.iter_rows(values_only=True)]
    en_tetes = ["" if h is None else str(h).replace("﻿", "").strip() for h in lignes[0]]
    donnees = [list(r) + [None] * (len(en_tetes) - len(r)) for r in lignes[1:]]
    return en_tetes, [dict(zip(en_tetes, r)) for r in donnees]


def vers_texte(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).strip()
    return t or None


def vers_date(v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=v)).date()
    t = str(v).strip()
    if not t:
        return None
    for fmt in ("%d/%m/%Y", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S",
                "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    return None


def vers_nombre(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    t = re.sub("[   ]", "", str(v).strip())
    if not t:
        return None
    try:
        return float(t.replace(",", "."))
    except ValueError:
        return None


# ---------------------------------------------------------------- OT_Detail

def ot_detail(chemin, ref_activites, ref_sites, longueur=9, separateur="-", feuille="",
              fichier_source=None, importe_le=None):
    """ref_activites : liste de codes ; ref_sites : liste de couples (site, équipe)."""
    en_tetes, lignes = lire_export(chemin, feuille)
    manquantes = [c for c in OBLIGATOIRES if c not in en_tetes]
    if manquantes:
        raise ErreurImport("Colonne(s) introuvable(s) dans l'extraction : " + ", ".join(manquantes))
    absentes = [c for c in ATTENDUES if c not in en_tetes]
    ignorees = [c for c in en_tetes if c not in ATTENDUES]
    ecarts = " | ".join(x for x in (
        "Colonnes absentes (laissées vides) : " + ", ".join(absentes) if absentes else "",
        "Colonnes non utilisées : " + ", ".join(ignorees) if ignorees else "") if x) or None

    codes = [str(c).strip().upper() for c in ref_activites if c is not None]
    sites = {str(s).strip() for s, _ in ref_sites if s is not None}
    couples = {f"{str(s).strip()}|{str(e).strip()}" for s, e in ref_sites if s is not None and e is not None}

    rows = []
    for n, brut in enumerate(lignes, start=2):
        r = {}
        for c in ATTENDUES:
            v = brut.get(c)
            r[c] = vers_date(v) if c in COLONNES_DATES else vers_nombre(v) if c in COLONNES_NOMBRES else vers_texte(v)
        r["Ligne export"] = n
        r["_DateBrute"] = vers_texte(brut.get("Date début planifiée"))
        pt = r["Poste technique"]
        r["Code poste"] = None if pt is None else pt[:longueur]
        r["Année planifiée"] = None if r["Date début planifiée"] is None else r["Date début planifiée"].year
        r["_CleOrdre"] = f"#LIGNE-{n}" if r["Ordre"] is None else r["Ordre"]
        rows.append(r)

    nb_par_ordre = {}
    for r in rows:
        nb_par_ordre[r["_CleOrdre"]] = nb_par_ordre.get(r["_CleOrdre"], 0) + 1

    eclate = []
    for r in rows:
        r["Nb lignes export de l'OT"] = nb_par_ordre[r["_CleOrdre"]]
        segments = [] if r["Activité"] is None else [s.strip().upper() for s in r["Activité"].split("+")]
        pt, site, equipe = r["Poste technique"], r["Groupe gestionnaires"], r["Poste travail princ."]
        ctl = [
            "Ordre vide" if r["Ordre"] is None else None,
            "OT présent sur plusieurs lignes de l'export"
            if r["Nb lignes export de l'OT"] > 1 and r["Ordre"] is not None else None,
            "Poste technique vide" if pt is None else None,
            "Poste technique trop court" if pt is not None and len(pt) < longueur else None,
            "Format du Poste technique inattendu"
            if pt is not None and separateur and len(pt) > longueur and pt[longueur] != separateur else None,
            "Date début planifiée vide" if r["_DateBrute"] is None else None,
            "Date début planifiée illisible"
            if r["_DateBrute"] is not None and r["Date début planifiée"] is None else None,
            "Activité vide" if r["Activité"] is None else None,
            "Activité mal formée" if r["Activité"] is not None and "" in segments else None,
            "Site ou équipe vide" if site is None or equipe is None else None,
            "Nouveau site" if site is not None and site not in sites else None,
            "Nouvelle équipe" if site is not None and equipe is not None and site in sites
            and f"{site}|{equipe}" not in couples else None,
            "Type de travail autre que PRE"
            if r["Type de travail"] is not None and r["Type de travail"] != "PRE" else None,
        ]
        r["_ControlesLigne"] = [c for c in ctl if c]
        activites = [s for s in segments if s != ""] or [None]
        for a in activites:
            e = dict(r)
            e["Activité atomique"] = a
            e["Ordre d'affichage activité"] = codes.index(a) + 1 if a in codes else 999
            e["_CleCouple"] = r["_CleOrdre"] + "|" + ("#VIDE" if a is None else a)
            eclate.append(e)

    occ, lignes_couple = {}, {}
    for e in eclate:
        occ[e["_CleCouple"]] = occ.get(e["_CleCouple"], 0) + 1
        lignes_couple.setdefault(e["_CleCouple"], set()).add(e["Ligne export"])
    vus, detail = set(), []
    for e in eclate:
        if e["_CleCouple"] in vus:
            continue
        vus.add(e["_CleCouple"])
        e["Occurrences Ordre+Activité"] = occ[e["_CleCouple"]]
        e["Lignes export regroupées"] = (", ".join(str(x) for x in sorted(lignes_couple[e["_CleCouple"]]))
                                         if occ[e["_CleCouple"]] > 1 else None)
        liste = list(e["_ControlesLigne"])
        if e["Activité atomique"] is not None and e["Ordre d'affichage activité"] == 999:
            liste.append("Activité inconnue du référentiel")
        if e["Occurrences Ordre+Activité"] > 1:
            liste.append("Couple Ordre + Activité en double")
        e["Contrôle qualité"] = " ; ".join(liste) or None
        e["Écarts de structure"] = ecarts
        e["Fichier source"] = fichier_source if fichier_source is not None else str(chemin)
        e["Importé le"] = importe_le or dt.datetime.now().replace(microsecond=0)
        detail.append({c: e[c] for c in COLONNES_DETAIL})
    return detail


# ---------------------------------------------------------------- OT_Couverture

def ot_couverture(detail):
    vide = lambda v: "(vide)" if v is None else v
    lignes = []
    for d in detail:
        lignes.append({
            "Année planifiée": d["Année planifiée"], "Site": vide(d["Groupe gestionnaires"]),
            "Équipe": vide(d["Poste travail princ."]), "Code poste": vide(d["Code poste"]),
            "Activité atomique": vide(d["Activité atomique"]), "Rang": d["Ordre d'affichage activité"],
            "_OT": f"#LIGNE-{d['Ligne export']}" if d["Ordre"] is None else d["Ordre"],
            "Contrôle qualité": d["Contrôle qualité"]})
    activites = []
    for a, _ in sorted({(l["Activité atomique"], l["Rang"]) for l in lignes}, key=lambda x: (x[1], x[0])):
        if a not in activites:
            activites.append(a)
    groupes = {}
    for l in lignes:
        groupes.setdefault((l["Année planifiée"], l["Site"], l["Équipe"], l["Code poste"]), []).append(l)
    resultat = []
    for (annee, site, equipe, poste), g in groupes.items():
        ligne = {"Année planifiée": annee, "Site": site, "Équipe": equipe, "Code poste": poste}
        for a in activites:
            ligne[a] = len({l["_OT"] for l in g if l["Activité atomique"] == a})
        ligne["OT distincts (toutes activités)"] = len({l["_OT"] for l in g})
        ligne["OT avec point de contrôle"] = len({l["_OT"] for l in g if l["Contrôle qualité"] is not None})
        resultat.append(ligne)
    resultat.sort(key=lambda l: (l["Année planifiée"] is not None, l["Année planifiée"] or 0,
                                 l["Site"], l["Équipe"], l["Code poste"]))
    colonnes = ["Année planifiée", "Site", "Équipe", "Code poste"] + activites + [
        "OT distincts (toutes activités)", "OT avec point de contrôle"]
    return colonnes, resultat


# ---------------------------------------------------------------- OT_Controles

def numeros_lignes(lignes_detail):
    """Lignes de l'export portées par des lignes du détail, doublons regroupés compris."""
    numeros = []
    for d in lignes_detail:
        numeros.append(d["Ligne export"])
        if d["Lignes export regroupées"]:
            numeros += [int(x) for x in d["Lignes export regroupées"].split(", ")]
    return list(dict.fromkeys(numeros))


def ot_controles(detail):
    lignes_export = numeros_lignes(detail)
    ecarts = sorted({d["Écarts de structure"] for d in detail if d["Écarts de structure"]})
    resultat = [
        {"N°": 0, "Contrôle": "Lignes lues dans l'export", "Niveau": "Info",
         "Lignes de l'export": len(lignes_export),
         "OT distincts": len({d["Ordre"] for d in detail if d["Ordre"] is not None}),
         "Exemples (Ordre ou ligne)": "", "Statut": "INFO"},
        {"N°": 1, "Contrôle": "Écart de structure de l'export", "Niveau": "À vérifier",
         "Lignes de l'export": 0, "OT distincts": 0, "Exemples (Ordre ou ligne)": " | ".join(ecarts),
         "Statut": "À VÉRIFIER" if ecarts else "OK"},
    ]
    for i, (libelle, niveau) in enumerate(CONTROLES):
        touchees = [d for d in detail if d["Contrôle qualité"] and libelle in d["Contrôle qualité"].split(" ; ")]
        num = sorted(numeros_lignes(touchees))
        ordres = list(dict.fromkeys(d["Ordre"] for d in touchees if d["Ordre"] is not None))
        exemples = ordres if ordres else [f"ligne {n}" for n in num]
        resultat.append({
            "N°": i + 2, "Contrôle": libelle, "Niveau": niveau, "Lignes de l'export": len(num),
            "OT distincts": len(ordres), "Exemples (Ordre ou ligne)": ", ".join(exemples[:10]),
            "Statut": "OK" if not num else "INFO" if niveau == "Info" else "À VÉRIFIER"})
    return resultat
